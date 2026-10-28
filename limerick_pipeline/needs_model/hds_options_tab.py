"""Add an 'Opportunity sites & options' tab to the HDS needs model, after 'Zoned land vs need'.

Per catchment: the shortfall from 'Zoned land vs need' (live), what the opportunity sites could add (values from
Limerick_Opportunity_Sites.gpkg), what is left, and the infill land that would take at Council density and at a higher
city infill density (lever). Catchments where infill is constrained (lever column) are shown separately so their remainder
can be placed elsewhere. Run after hds_land_tab.py: python hds_options_tab.py <in.xlsx> <out.xlsx>
"""
import sys
import pandas as pd, geopandas as gpd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

SRC, OUT = sys.argv[1], sys.argv[2]
g = gpd.read_file('Limerick_Opportunity_Sites.gpkg', layer='opportunity_sites')
g['ok'] = g.residential_permitted_in_zone.isin(['Yes', 'Mixed use'])
opp = g.groupby('catchment').apply(lambda d: pd.Series(dict(
    sites=len(d), ha=d.area_ha.sum(), ok=d.loc[d.ok, 'indicative_homes_not_in_register'].sum(),
    rezone=d.loc[~d.ok, 'indicative_homes_not_in_register'].sum())))
rail = opp.loc['Rail corridor (long term)'] if 'Rail corridor (long term)' in opp.index else None
CONSTRAINED = {'Kings Island': 'Severe flood constraint across most of the island (analyst)'}

wb = openpyxl.load_workbook(SRC)
Z = wb['Zoned land vs need']
NAME = 'Opportunity sites & options'
if NAME in wb.sheetnames:
    del wb[NAME]
ws = wb.create_sheet(NAME, index=wb.sheetnames.index('Zoned land vs need') + 1)

F = 'Arial'
DARK, MUTED, BAND = 'FF0E2925', 'FF3E5853', 'FFF4EFE8'
hdr_font, hdr_fill = Font(name=F, size=10, bold=True, color='FFFFFFFF'), PatternFill('solid', fgColor=DARK)
body, bold, note = Font(name=F, size=10), Font(name=F, size=10, bold=True), Font(name=F, size=10, color=MUTED)
lever_font, lever_fill = Font(name=F, size=10, bold=True, color='FF0000FF'), PatternFill('solid', fgColor='FFE4E66A')
band = PatternFill('solid', fgColor=BAND)
INT, DEC = '#,##0;\\(#,##0\\);\\-', '#,##0.0;\\(#,##0.0\\);\\-'
h2 = Font(name=F, size=12, bold=True, color=DARK)

ws['A1'] = 'Opportunity sites and options for the 2040 shortfall'; ws['A1'].font = Font(name=F, size=16, bold=True, color=DARK)
ws['A2'] = ('Starts from the shortfall in the Zoned land vs need tab, adds the opportunity sites identified on the desk review, and shows '
            'how much infill land would still be needed. Opportunity-site homes are indicative (full site at Council density) and are not supply.')
ws['A2'].font = note
ws['A4'] = 'City infill density (homes/ha)'; ws['A4'].font = bold
ws['C4'] = 100; ws['C4'].font, ws['C4'].fill = lever_font, lever_fill
ws['D4'] = 'Density assumed for the remaining city infill on sites not yet identified. Change to test; Variation No. 3 allows 50–200 in city Zone 2 and 100–250 in the city centre.'
ws['D4'].font = note
ws['A5'] = 'Headroom'; ws['A5'].font = bold
ws['C5'] = "='Zoned land vs need'!C4"; ws['C5'].number_format = '0%'; ws['C5'].font = body
ws['D5'] = 'Taken from the Zoned land vs need tab.'; ws['D5'].font = note

heads = ['Catchment', 'Type of place', 'Council density (homes/ha)', 'Homes needed 2026–2040', 'Identified supply (homes)',
         'Shortfall now (homes)', 'Opportunity sites', 'Opportunity homes: zoning allows housing', 'Opportunity homes: needs rezoning',
         'Left after sites where zoning allows', 'Left after all opportunity sites', 'Infill constraint (enter to flag)',
         'Left in unconstrained catchments', 'Left in constrained catchments (to place elsewhere)',
         'Infill land still needed at Council density (ha)', 'Infill land still needed at city infill density (ha)',
         'Left after all sites, with headroom (homes)']
HR = 7
for j, h in enumerate(heads, 1):
    c = ws.cell(HR, j, h); c.font, c.fill = hdr_font, hdr_fill; c.alignment = Alignment(wrap_text=True, vertical='top')
ws.cell(HR - 1, 4, 'From Zoned land vs need').font = bold
ws.cell(HR - 1, 7, 'From the opportunity sites register').font = bold
ws.cell(HR - 1, 12, 'Analyst input').font = bold

r0 = HR + 1
for k in range(24):
    r, z = r0 + k, 7 + k
    name = wb['Land at Council density'].cell(z, 1).value         # plain names (the Zoned land tab holds formulas)
    o = opp.loc[name] if name in opp.index else None
    f = {1: f"='Zoned land vs need'!A{z}", 2: f"='Zoned land vs need'!B{z}", 3: f"='Zoned land vs need'!C{z}",
         4: f"='Zoned land vs need'!D{z}", 5: f"='Zoned land vs need'!I{z}", 6: f'=MAX(D{r}-E{r},0)',
         7: int(o.sites) if o is not None else 0, 8: round(float(o.ok), 0) if o is not None else 0,
         9: round(float(o.rezone), 0) if o is not None else 0,
         10: f'=MAX(F{r}-H{r},0)', 11: f'=MAX(J{r}-I{r},0)', 12: CONSTRAINED.get(name, ''),
         13: f'=IF(L{r}="",K{r},0)', 14: f'=IF(L{r}="",0,K{r})',
         15: f'=IF(C{r}>0,K{r}/C{r},0)', 16: f'=IF(B{r}="City neighbourhood",IF($C$4>0,K{r}/MAX($C$4,C{r}),0),IF(C{r}>0,K{r}/C{r},0))',
         17: f'=MAX(D{r}*(1+$C$5)-E{r}-H{r}-I{r},0)'}
    for j, v in f.items():
        c = ws.cell(r, j, v); c.font = bold if j in (11, 16) else body
        c.number_format = {3: '0.0', 15: DEC, 16: DEC}.get(j, 'General' if j in (1, 2, 12) else INT)
        if j == 12:
            c.fill = lever_fill; c.font = Font(name=F, size=10, color='FF0000FF')
rt = r0 + 24
ws.cell(rt, 1, 'Neighbourhoods and settlements')
for j in range(1, 18):
    ws.cell(rt, j).fill = band; ws.cell(rt, j).font = bold
for j in [4, 5, 6, 7, 8, 9, 10, 11, 13, 14, 15, 16, 17]:
    L = openpyxl.utils.get_column_letter(j)
    c = ws.cell(rt, j, f'=SUM({L}{r0}:{L}{rt - 1})'); c.number_format = ws.cell(r0, j).number_format
if rail is not None:
    rr = rt + 1
    for j, v in {1: 'Rail corridor new towns (long term, not counted)', 7: int(rail.sites), 9: round(float(rail.rezone), 0)}.items():
        c = ws.cell(rr, j, v); c.font = note; c.number_format = INT if j != 1 else 'General'
    ws.cell(rr, 10, 'Beyond 2040; shown for completeness').font = note

# by type of place
tr = rt + 3
ws.cell(tr, 1, 'By type of place').font = h2
th = ['Type of place', '', '', 'Homes needed', 'Identified supply', 'Shortfall now', 'Opportunity sites', 'Zoning allows', 'Needs rezoning',
      'Left after zoning-allows sites', 'Left after all sites', 'Catchments still short', 'Left, unconstrained', 'Left, constrained',
      'Infill ha at Council density', 'Infill ha at city infill density', 'Left with headroom']
for j, h in enumerate(th, 1):
    c = ws.cell(tr + 1, j, h); c.font, c.fill = hdr_font, hdr_fill; c.alignment = Alignment(wrap_text=True, vertical='top')
for k, t in enumerate(['City neighbourhood', 'Key town', 'Village']):
    r = tr + 2 + k
    ws.cell(r, 1, t).font = body
    for j in [4, 5, 6, 7, 8, 9, 10, 11, 13, 14, 15, 16, 17]:
        L = openpyxl.utils.get_column_letter(j)
        c = ws.cell(r, j, f'=SUMIFS({L}${r0}:{L}${rt - 1},$B${r0}:$B${rt - 1},$A{r})'); c.font = body
        c.number_format = ws.cell(r0, j).number_format
    ws.cell(r, 12, f'=COUNTIFS($B${r0}:$B${rt - 1},$A{r},$K${r0}:$K${rt - 1},">0")').font = body

# reading of the results (live sentences)
sr = tr + 7
ws.cell(sr, 1, 'What the numbers say').font = h2
city, town, vill = tr + 2, tr + 3, tr + 4
lines = [
    f'="Key towns: "&TEXT(F{town},"#,##0")&" homes short now across "&COUNTIFS($B${r0}:$B${rt-1},"Key town",$F${r0}:$F${rt-1},">0")&" town(s); after opportunity sites, "&TEXT(K{town},"#,##0")&" short."',
    f'="Villages: "&TEXT(F{vill},"#,##0")&" homes short now across "&COUNTIFS($B${r0}:$B${rt-1},"Village",$F${r0}:$F${rt-1},">0")&" villages; after opportunity sites where zoning allows, "&TEXT(J{vill},"#,##0")&"; after all opportunity sites, "&TEXT(K{vill},"#,##0")&"."',
    f'="City neighbourhoods: "&TEXT(F{city},"#,##0")&" homes short now. Opportunity sites where zoning allows cover "&TEXT(F{city}-J{city},"#,##0")&"; with rezoned opportunity sites, "&TEXT(F{city}-K{city},"#,##0")&". That leaves "&TEXT(K{city},"#,##0")&" homes for infill on sites not yet identified."',
    f'="That remaining city infill needs about "&TEXT(O{city},"#,##0")&" ha at Council densities, or about "&TEXT(P{city},"#,##0")&" ha at "&$C$4&" homes/ha."',
    f'="Constrained catchments (flagged in column L) leave "&TEXT(N{rt},"#,##0")&" homes to be placed in other neighbourhoods."',
]
for i, t in enumerate(lines):
    c = ws.cell(sr + 1 + i, 1, t); c.font = body

orow = sr + 8
ws.cell(orow, 1, 'Options for closing the shortfall').font = h2
opts = [
    ('1. Bring forward opportunity sites where zoning already allows housing',
     'Column H. No plan change is needed; the work is ownership, servicing and delivery route. Confirm each site in the opportunity sites register first, because the homes are indicative.'),
    ('2. Rezone opportunity sites through the plan review',
     'Column I. Most of these homes sit on a few large city sites zoned open space, agriculture or education (Caherdavin, Southhill, Dooradoyle). Check whether those are true infill before relying on them.'),
    ('3. Higher density on city infill not yet identified',
     'Columns O–P and the lever in C4. The city shortfall left after the identified land and opportunity sites has to come from smaller infill and brownfield sites at well above current densities.'),
    ('4. Vacancy, dereliction and over-the-shop',
     'Not counted here. Bringing vacant and derelict buildings back into use adds homes without new land, especially in the City Centre and town centres. The register’s GeoDirectory counts cover residential addresses only, so this needs its own survey.'),
    ('5. Place constrained neighbourhoods’ share elsewhere',
     'Column N. Where infill is constrained (e.g. Kings Island, flood risk across most of the island), the remaining homes should be planned for in neighbouring neighbourhoods or the City Centre.'),
    ('6. Rail corridor new towns (after 2040)',
     'Shown below the total and not counted. A longer-term option along the rail lines.'),
    ('Settlements outside the city',
     'Key towns already have more zoned land than they need, except Abbeyfeale (small). Several villages are short by a few hectares; their opportunity sites close part of the gap. See the sentences above.'),
]
for i, (a, b) in enumerate(opts):
    ws.cell(orow + 1 + i, 1, a).font = bold
    ws.cell(orow + 1 + i, 4, b).font = note

nrow = orow + 10
ws.cell(nrow, 1, 'Notes').font = h2
notes = [
    'Opportunity-site homes: site area × 0.91 × the Council density for the catchment, less any part already in the land register. They are an upper bound until each site is investigated; confidence is Low.',
    '“Zoning allows housing” = the main zoning on the site is residential or mixed use (town or village centre, district or local centre). Everything else needs rezoning.',
    'Infill constraint: enter a short reason in column L to flag a catchment; its remaining homes then move to column N. Kings Island is flagged by the analyst for flood risk.',
    'Infill land is net residential land. “City infill density” applies only to city neighbourhoods and never goes below the Council density there.',
]
for i, t in enumerate(notes):
    ws.cell(nrow + 1 + i, 1, t).font = note

widths = [30, 18, 11, 11, 11, 11, 10, 12, 12, 12, 12, 26, 12, 13, 12, 12, 12]
for j, w in enumerate(widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
ws.row_dimensions[HR].height = 66; ws.row_dimensions[tr + 1].height = 42
ws.freeze_panes = 'C8'
I = wb['Intro']
for r in range(31, 45):
    if I.cell(r, 2).value == 'Zoned land vs need':
        I.insert_rows(r + 1)
        I.cell(r + 1, 2, NAME).font = I.cell(r, 2).font.copy()
        I.cell(r + 1, 4, 'What the opportunity sites could add to the shortfall, what is left for infill, and options to close it (city infill density lever).').font = I.cell(r, 4).font.copy()
        break
wb.save(OUT)
print('saved', OUT)
