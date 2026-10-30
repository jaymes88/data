"""Delivery estimates should cover only homes not yet built.

For each permission, the yearly estimates (est_q4_2026 ... est_2031_plus) can add up to at most units − units_built.
Any excess comes off the earliest years first. units_built is all units for a completed scheme, otherwise GeoDirectory's
built count, so this nets part-built schemes the same way future supply does. The rule is idempotent.

Used by build_land_register.py. Run on its own to apply it to an existing register:
    python net_built.py [Limerick_Housing_Land_Register.gpkg] [Limerick_Housing_Land_Register.xlsx]
"""
import sys, sqlite3
import numpy as np, pandas as pd


def net_estimates(pp, yc):
    pp = pp.copy()
    left = (pp.units.fillna(0) - pp.units_built.fillna(0)).clip(lower=0)
    excess = (pp[yc].fillna(0).sum(axis=1) - left).clip(lower=0)
    excess = excess.where(excess > 0.05, 0)        # ignore rounding in the source estimates
    for c in yc:                                   # earliest year first
        take = np.minimum(pp[c].fillna(0), excess)
        pp[c] = pp[c].fillna(0) - take
        excess = excess - take
    return pp


if __name__ == '__main__':
    REG = sys.argv[1] if len(sys.argv) > 1 else 'Limerick_Housing_Land_Register.gpkg'
    XL = sys.argv[2] if len(sys.argv) > 2 else 'Limerick_Housing_Land_Register.xlsx'
    YC = ['est_q4_2026', 'est_2027', 'est_2028', 'est_2029', 'est_2030', 'est_2031_plus']
    con = sqlite3.connect(REG)
    pp = pd.read_sql(f"select fid, perm_id, parcel_id, units, units_built, status, counts_toward_totals, {', '.join(YC)} from parcel_permissions", con)
    nw = net_estimates(pp, YC)
    ch = nw[(nw[YC] - pp[YC].fillna(0)).abs().sum(axis=1) > 0.05]
    for _, r in ch.iterrows():
        con.execute(f"update parcel_permissions set {', '.join(c + ' = ?' for c in YC)} where fid = ?", [float(r[c]) for c in YC] + [int(r.fid)])
    con.commit()
    cnt = pp.counts_toward_totals.eq('Yes')
    print(f'{len(ch)} permissions netted; counted estimates {pp.loc[cnt, YC].sum().sum():,.1f} -> {nw.loc[cnt, YC].sum().sum():,.1f}')
    # same values into the register workbook's Permissions sheet (Parcels and Catchments sum them with formulas)
    import openpyxl
    wb = openpyxl.load_workbook(XL)
    ws = wb['Permissions']
    head = {c.value: c.column for c in ws[1]}
    new = nw.set_index(nw.perm_id.astype(str))
    n = 0
    for r in range(2, ws.max_row + 1):
        pid = str(ws.cell(r, head['perm_id']).value)
        if pid in new.index and pid in set(ch.perm_id.astype(str)):
            row = new.loc[pid]
            row = row.iloc[0] if isinstance(row, pd.DataFrame) else row
            for c in YC:
                ws.cell(r, head[c]).value = float(row[c]); n += 1
    wb.save(XL)
    print(f'{XL}: {n} cells updated')
