"""Status overrides from later evidence (inputs/status_overrides.csv), e.g. a building control commencement notice for a
scheme the register still shows as not started.

Each listed permission takes the new status and start date, and its delivery estimates are re-projected with the supply
dashboard's method for schemes on site (supply/supply_dashboard.py, allocate()): build time = the larger of the median
start-to-completion time for its size band and units / median build rate; first homes a year after start (or halfway
through a short build); units spread evenly to the end of the build, counting only time after 08/10/2026. Estimates are
then netted of homes already built (net_built.py).

Used by build_land_register.py. Run on its own to apply to an existing register:
    python status_overrides.py [register.gpkg] [register.xlsx] [supply.gpkg]
"""
import os, sys, sqlite3
import numpy as np, pandas as pd, geopandas as gpd

YC = ['est_q4_2026', 'est_2027', 'est_2028', 'est_2029', 'est_2030', 'est_2031_plus']
YB = [(2026.0, 2027.0), (2027.0, 2028.0), (2028.0, 2029.0), (2029.0, 2030.0), (2030.0, 2031.0), (2031.0, 9999.0)]
T0 = pd.Timestamp('2026-10-08')
yr = lambda t: t.year + (t.day_of_year - 1) / (366 if t.is_leap_year else 365)
NOW = yr(T0)
q = lambda s: f'{int(s)} Q{min(4, int(round((s % 1) * 4, 6)) + 1)}'


def band_params(supply_gpkg):
    P = gpd.read_file(supply_gpkg, layer='permissions', ignore_geometry=True)
    par = {b: dict(dur=g.start_to_complete_yrs.median(), rate=g.build_rate_dpa[g.build_rate_dpa > 0].median()) for b, g in P.groupby('size_band')}
    par['Unknown'] = dict(dur=P.start_to_complete_yrs.median(), rate=P.build_rate_dpa[P.build_rate_dpa > 0].median())
    return par, P.set_index(P.perm_id.astype(str)).size_band


def on_site(units, start, band, par):
    """Delivery estimates for a scheme on site (same as allocate() in supply_dashboard.py)."""
    p = par.get(band, par['Unknown'])
    dur = max(p['dur'], units / p['rate'])
    s = yr(pd.Timestamp(start))
    first, end = s + min(1.0, dur / 2), s + dur
    a = max(first, NOW)
    if end - a < 1.0:
        end = a + 1.0
    est = [round(units * max(0.0, min(end, hi) - max(a, lo)) / (end - a), 1) for lo, hi in YB]
    est[-1] = round(est[-1] + units - sum(est), 1)          # rounding residue to the last year, as the dashboard does
    return est, q(s), q(end)


def apply(pp, overrides, par, bands):
    pp = pp.copy()
    for _, o in overrides.iterrows():
        m = pp.planning_ref.astype(str).eq(str(o.planning_ref))
        assert m.any(), ('override not in register', o.planning_ref)
        for i in pp.index[m]:
            pp.at[i, 'status'] = o.status
            pp.at[i, 'status_source'] = o.source
            pp.at[i, 'start_on_site'] = o.start_on_site
            if o.status == 'Under construction':
                est, s, e = on_site(float(pp.at[i, 'units']), o.start_on_site, bands.get(str(pp.at[i, 'perm_id']), 'Unknown'), par)
                for c, v in zip(YC, est):
                    pp.at[i, c] = v
                pp.at[i, 'est_start'], pp.at[i, 'est_last_completions'], pp.at[i, 'confidence'] = s, e, 'High'
    from net_built import net_estimates
    return net_estimates(pp, YC)


if __name__ == '__main__':
    REG = sys.argv[1] if len(sys.argv) > 1 else 'Limerick_Housing_Land_Register.gpkg'
    XL = sys.argv[2] if len(sys.argv) > 2 else 'Limerick_Housing_Land_Register.xlsx'
    SUP = sys.argv[3] if len(sys.argv) > 3 else '../supply/Limerick_Housing_Supply.gpkg'
    OV = pd.read_csv('inputs/status_overrides.csv', dtype=str)
    par, bands = band_params(SUP)
    F = ['status', 'status_source', 'start_on_site', 'est_start', 'est_last_completions', 'confidence'] + YC
    con = sqlite3.connect(REG)
    pp = pd.read_sql(f"select fid, perm_id, planning_ref, units, units_built, {', '.join(F)} from parcel_permissions", con)
    nw = apply(pp, OV, par, bands)
    same = lambda a, b: (a.isna() & b.isna()) | (a.astype(str) == b.astype(str)) | (pd.to_numeric(a, errors='coerce').sub(pd.to_numeric(b, errors='coerce')).abs() < 0.05)
    ch = nw[~pd.concat([same(nw[c], pp[c]) for c in F], axis=1).all(axis=1)]
    for _, r in ch.iterrows():
        con.execute(f"update parcel_permissions set {', '.join(c + ' = ?' for c in F)} where fid = ?",
                    [None if pd.isna(r[c]) else (float(r[c]) if c in YC else str(r[c])) for c in F] + [int(r.fid)])
    con.commit()
    print(ch[['planning_ref', 'units', 'status', 'start_on_site', 'est_start', 'est_last_completions'] + YC].to_string())
    import openpyxl
    wb = openpyxl.load_workbook(XL)
    ws = wb['Permissions']
    head = {c.value: c.column for c in ws[1]}
    new = nw.set_index(nw.perm_id.astype(str))
    n = 0
    for r in range(2, ws.max_row + 1):
        pid = str(ws.cell(r, head['perm_id']).value)
        if pid in set(ch.perm_id.astype(str)):
            for c in F:
                v = new.at[pid, c]
                ws.cell(r, head[c]).value = None if pd.isna(v) else (float(v) if c in YC else v)
                n += 1
    wb.save(XL)
    print(f'{len(ch)} permissions changed; {XL}: {n} cells updated')
