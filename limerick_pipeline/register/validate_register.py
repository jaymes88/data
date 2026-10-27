"""Validation suite for the Limerick housing land register and the files that feed it.

Three groups of checks, each PASS / WARN / FAIL with the rows behind it:
  I  integrity      – keys, links, overlaps, geometry, unit conservation, cross-file agreement, plausibility
  P  provenance     – every counted fact traced back to its source: live national planning register, appeal cases,
                      building control, and each hand-entered input (reinstated sites, added schemes, unit mixes)
  E  external       – CSO new dwelling completions and planning permissions against the register and GeoDirectory
plus a stratified QA sample for a person to check against the planning files.

Re-run after every rebuild, and after the county settlements audit is added (checks are scope-agnostic; C-checks
flag county audit parcels that overlap interim zoned-land estimates).
Writes Validation_Report.xlsx.
"""
import json, re, os, sqlite3, time
import numpy as np, pandas as pd, geopandas as gpd, requests
from shapely.ops import unary_union
from shapely.geometry import shape

ITM = 2157
REG = 'Limerick_Housing_Land_Register.gpkg'
SUP = 'Limerick_Housing_Supply.gpkg'
WB = 'Limerick_Residential_Pipeline_Simplified.xlsx'
DASH = 'Limerick_Housing_Supply_Dashboard.xlsx'
OUT = 'Validation_Report.xlsx'
RUN = pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')
results, detail = [], {}


def check(cid, group, name, status, n, note, rows=None):
    results.append(dict(check=cid, group=group, name=name, status=status, items=int(n), detail=note))
    if rows is not None and len(rows):
        detail[cid] = rows.copy()
    print(f'{cid:4} {status:4} {name} – {note}')


con = sqlite3.connect(REG)
parcels = gpd.read_file(REG, layer='parcels').to_crs(ITM)
ps = gpd.read_file(REG, layer='parcel_summary').to_crs(ITM)
pp = pd.read_sql('select * from parcel_permissions', con)
for c in ps.columns:
    if c not in ('geometry',) and ps[c].dtype == object:
        try:
            ps[c] = pd.to_numeric(ps[c])
        except (ValueError, TypeError):
            pass
cnt = pp[pp.counts_toward_totals.eq('Yes')]
LIVE = ['Complete', 'Under construction', 'Permitted – not started', 'Stalled']

# ======================================================================== I: integrity
d = parcels[parcels.parcel_id.duplicated(keep=False)]
check('I1', 'Integrity', 'parcel_id unique', 'PASS' if d.empty else 'FAIL', len(d), f'{len(parcels)} parcels', d[['parcel_id']])
d = pp[pp.perm_id.duplicated(keep=False)]
check('I2', 'Integrity', 'perm_id unique', 'PASS' if d.empty else 'FAIL', len(d), f'{len(pp)} permissions', d[['perm_id']])
d = pp[~pp.parcel_id.isin(parcels.parcel_id)]
check('I3', 'Integrity', 'every permission links to an existing parcel', 'PASS' if d.empty else 'FAIL', len(d), 'one permission, one parcel', d[['perm_id', 'parcel_id']])
r = cnt[cnt.planning_ref.ne('')]
d = r[r.planning_ref.duplicated(keep=False)]
check('I4', 'Integrity', 'no planning reference counted twice', 'PASS' if d.empty else 'FAIL', len(d), 'counted rows only', d[['perm_id', 'planning_ref', 'units']])

ov = gpd.sjoin(parcels[['parcel_id', 'geometry']], parcels[['parcel_id', 'geometry']], predicate='intersects')
ov = ov[ov.parcel_id_left < ov.parcel_id_right]
g = parcels.set_index('parcel_id').geometry
ov['overlap_ha'] = [round(g[a].intersection(g[b]).area / 1e4, 4) for a, b in zip(ov.parcel_id_left, ov.parcel_id_right)]
ov = ov[ov.overlap_ha > 0.01]
check('I5', 'Integrity', 'parcels do not overlap (>0.01 ha)', 'PASS' if ov.empty else 'FAIL', len(ov), f'{ov.overlap_ha.sum():.2f} ha overlapping',
      pd.DataFrame(ov[['parcel_id_left', 'parcel_id_right', 'overlap_ha']]))

bad = parcels[~parcels.is_valid | parcels.is_empty]
fc = lambda p: gpd.GeoDataFrame(geometry=[shape(f['geometry']) for f in json.load(open(p))['features']], crs=4326).to_crs(ITM)
county = unary_union(list(fc('lea.geojson').geometry)).buffer(200)
outside = parcels[~parcels.representative_point().within(county)]
check('I6', 'Integrity', 'geometry valid and inside Limerick', 'PASS' if bad.empty and outside.empty else 'FAIL', len(bad) + len(outside),
      f'{len(bad)} invalid/empty, {len(outside)} outside the county', pd.concat([bad[['parcel_id']].assign(issue='invalid'), outside[['parcel_id']].assign(issue='outside county')]))

buckets = ['complete_units', 'under_construction_units', 'permitted_not_started_units', 'stalled_units', 'expired_units', 'in_planning_units']
a, b = cnt.units.sum(), ps[buckets].sum().sum()
c1 = abs(ps.total_supply_units - ps.consented_units - ps.unconsented_capacity).max()
check('I7', 'Integrity', 'unit conservation (permissions = parcel buckets; total = consented + unconsented)', 'PASS' if abs(a - b) < 0.5 and c1 < 0.5 else 'FAIL', 0,
      f'permission units {a:,.0f} vs parcel buckets {b:,.0f}; max parcel imbalance {c1:.2f}')

# cross-file agreement: tracker workbook -> supply permissions -> register; dashboard totals
sup = gpd.read_file(SUP, layer='permissions')
sup['units'] = sup.units.fillna(0)
m = sup[['perm_id', 'units', 'status']].merge(pp[['perm_id', 'units', 'status_cis', 'status']], on='perm_id', how='outer', suffixes=('_supply', '_register'), indicator=True)
# permissions the analyst removed at review are expected to be missing from the register
import os
rmv = set()
if os.path.exists('in_review/removed_parcels.gpkg'):
    rr = gpd.read_file('in_review/removed_parcels.gpkg')
    refs = {x for v in rr.planning_refs.fillna('') for x in v.split(';') if x}
    rmv = set(sup[sup.planning_ref.astype(str).isin(refs) | sup.register_ref.astype(str).isin(refs)].perm_id)
diff = m[((m._merge != 'both') | (m.units_supply.fillna(-1) != m.units_register.fillna(-1))) & ~m.perm_id.isin(rmv)]
check('I8', 'Integrity', 'supply GeoPackage and register hold the same permissions and units (less review removals)', 'PASS' if diff.empty else 'FAIL', len(diff),
      f'{len(sup)} vs {len(pp)} permissions; {len(rmv)} removed at review', diff)
wbk = pd.read_excel(WB, sheet_name='Applications', dtype={'Project Id': str})
wc = wbk[wbk['Counts?'].eq('Yes')][['Project Id', 'Units', 'Register Ref']].rename(columns={'Project Id': 'perm_id'})
m = wc.merge(pp[['perm_id', 'units', 'source']], on='perm_id', how='outer', indicator=True)
expl = m._merge.eq('right_only') & m.source.astype(str).str.startswith('Register only')        # register-only additions
dup8 = m._merge.eq('left_only')                                                                    # workbook rows not in register
d = m[(m._merge != 'both') & ~expl]
check('I9', 'Integrity', 'tracker workbook counted schemes all reach the register', 'PASS' if d.empty else 'WARN', len(d),
      f'{len(wc)} counted in workbook, {int(expl.sum())} register-only additions explained', d)
try:
    from openpyxl import load_workbook
    D = load_workbook(DASH, data_only=True)['Dashboard']
    kv = {D.cell(r, 1).value: D.cell(r, 2).value for r in range(5, 25)}
    dash = kv.get('All consented units')
    regc = ps.consented_units.sum()
    # the register moves schemes to Complete where GeoDirectory shows them built; the dashboard uses CIS status only.
    # Expired or in-planning records moved that way join the consented total in the register alone.
    rc = pp[pp.counts_toward_totals.eq('Yes') & pp.status.eq('Complete') & pp.status_cis.isin(['Expired – not started', 'In planning'])]
    adj = rc.units.sum()
    check('I10', 'Integrity', 'dashboard consented total = register consented total (after GeoDirectory completions)',
          'PASS' if dash and abs(dash + adj - regc) < 0.5 else 'FAIL', len(rc),
          f'dashboard {dash:,.0f} + {adj:,.0f} expired/in-planning units found built on GeoDirectory = register {regc:,.0f}',
          rc[['perm_id', 'parcel_id', 'planning_ref', 'units', 'status_cis', 'status', 'geod_new_built']])
except Exception as e:
    check('I10', 'Integrity', 'dashboard consented total = register consented total', 'WARN', 0, f'could not read dashboard: {e}')

ps['dph'] = ps.capacity_units / ps.developable_area_ha.replace(0, np.nan)
cap = ps[ps.capacity_units.notna() & ps.parcel_source.ne('Planning consent – not in council audit')]
d = cap[(cap.dph > 250) | (cap.dph < 5)][['parcel_id', 'parcel_source', 'local_name', 'developable_area_ha', 'capacity_units', 'dph', 'capacity_basis']]
check('I11', 'Integrity', 'capacity density plausible (5–250 homes/ha)', 'PASS' if d.empty else 'WARN', len(d), 'council and estimated capacities', d)
d = ps[ps.flag.eq('Consented above capacity')][['parcel_id', 'local_name', 'capacity_units', 'consented_units']]
check('I12', 'Integrity', 'consents above council capacity (capacity to review)', 'WARN' if len(d) else 'PASS', len(d), 'not an error: capacity estimates look low here', d)

d = ps[ps.parcel_source.eq('Planning consent – not in council audit') & ps.consented_units.gt(0) & ps.mapped_area_ha.gt(2)
       & (ps.density_estimate_units > 3 * ps.consented_units)][['parcel_id', 'local_name', 'mapped_area_ha', 'consented_units', 'density_estimate_units', 'parcel_status']]
check('I15', 'Integrity', 'consent parcel area fits its scheme (indicative yield under 3x consented units)', 'PASS' if d.empty else 'WARN', len(d),
      'parcel boundary may take in more land than the scheme, or the scheme is low density; check the outline', d.sort_values('mapped_area_ha', ascending=False))

for c in ['final_grant', 'expiry', 'start_on_site']:
    pp[c] = pd.to_datetime(pp[c], errors='coerce')
iss = []
iss += [('start before grant', r) for r in pp[pp.start_on_site < pp.final_grant - pd.Timedelta(days=30)].perm_id]
iss += [('expiry before grant', r) for r in pp[pp.expiry < pp.final_grant].perm_id]
iss += [('start in the future', r) for r in pp[pp.start_on_site > pd.Timestamp.now()].perm_id]
yc = [c for c in pp.columns if c.startswith('est_') and c not in ('est_start', 'est_last_completions')]
iss += [('complete but still projected', r) for r in pp[pp.status.eq('Complete') & pp[yc].sum(axis=1).gt(0)].perm_id]
iss += [('not live but projected', r) for r in pp[pp.status.isin(['Expired – not started', 'In planning']) & pp[yc].sum(axis=1).gt(0)].perm_id]
d = pd.DataFrame(iss, columns=['issue', 'perm_id']).merge(pp[['perm_id', 'planning_ref', 'scheme', 'status', 'final_grant', 'start_on_site', 'expiry']], on='perm_id')
check('I13', 'Integrity', 'dates and status consistent', 'PASS' if d.empty else 'WARN', len(d), 'start/expiry/grant order; projections only on live schemes', d)
d = ps[ps.catchment.isna() | (ps.catchment_share.fillna(0).eq(0) & ps.catchment_basis.eq('Inside catchment'))]
check('I14', 'Integrity', 'every parcel has a catchment and a basis', 'PASS' if d.empty else 'FAIL', len(d), ps.catchment_basis.str.replace(r' the .*', '', regex=True).value_counts().to_dict().__str__(), d[['parcel_id', 'catchment', 'catchment_basis']])

# ======================================================================== P: provenance
# P1/P2 live national planning register: does every counted reference exist, and is it still granted?
B = 'https://services.arcgis.com/NzlPQPKn5QF9v2US/arcgis/rest/services/IrishPlanningApplications/FeatureServer/0/query'
refs = sorted(set(r for r in pp.register_ref.fillna('').tolist() + pp.planning_ref.fillna('').tolist() if re.fullmatch(r'\d{4,8}', r)))
live = []
for i in range(0, len(refs), 80):
    q = "PlanningAuthority LIKE '%Limerick%' AND ApplicationNumber IN (" + ','.join(f"'{r}'" for r in refs[i:i + 80]) + ')'
    for t in range(3):
        try:
            j = requests.get(B, params={'where': q, 'outFields': 'ApplicationNumber,Decision,AppealDecision,ApplicationStatus,NumResidentialUnits,WithdrawnDate,ExpiryDate,DevelopmentDescription',
                                        'returnGeometry': 'false', 'f': 'json'}, timeout=120).json()
            live += [f['attributes'] for f in j.get('features', [])]
            break
        except Exception:
            time.sleep(3)
L = pd.DataFrame(live).drop_duplicates('ApplicationNumber') if live else pd.DataFrame(columns=['ApplicationNumber'])
L['ApplicationNumber'] = L.ApplicationNumber.astype(str).str.strip()
pp['ref'] = np.where(pp.register_ref.fillna('').str.fullmatch(r'\d{4,8}'), pp.register_ref, pp.planning_ref)
x = pp.merge(L, left_on='ref', right_on='ApplicationNumber', how='left')
nf = x[x.ApplicationNumber.isna() & x.ref.fillna('').str.fullmatch(r'\d{4,8}')].copy()
# the national dataset starts in 2017 and does not hold every Part 8 (council's own schemes, filed yy8nnn)
nf['reason'] = np.where(nf.ref.str.fullmatch(r'\d{2}8\d{3}'), 'Part 8 (council scheme) – not in national dataset',
                np.where(nf.ref.str[:2].astype(int) < 17, 'Pre-2017 file – before national dataset coverage', 'Not found – check'))
cov = nf[nf.reason.ne('Not found – check')]
nf_bad = nf[nf.reason.eq('Not found – check')]
check('P1', 'Provenance', 'every council reference exists in the live national planning register', 'PASS' if nf_bad.empty else 'FAIL', len(nf_bad),
      f'{len(L)} of {len(refs)} references found live ({pd.Timestamp.now():%d/%m/%Y}); {len(cov)} outside the dataset\'s coverage (Part 8 / pre-2017)', nf[['perm_id', 'ref', 'scheme', 'units', 'source', 'reason']])
noref = pp[~pp.ref.fillna('').str.fullmatch(r'\d{4,8}')]
check('P1b', 'Provenance', 'permissions without a council reference (CIS-only, Part 8 CIS ids, SHD)', 'WARN' if len(noref) else 'PASS', len(noref),
      'traced to CIS / appeal case instead', noref[['perm_id', 'planning_ref', 'scheme', 'units', 'source', 'status']])
dead = x[x.counts_toward_totals.eq('Yes') & x.status.isin(LIVE) & (
    x.Decision.fillna('').str.upper().str.contains('REFUS') & ~x.AppealDecision.fillna('').str.upper().str.contains('CONDITIONAL|GRANT')
    | x.AppealDecision.fillna('').str.upper().str.contains('REFUS') | x.WithdrawnDate.notna())]
check('P2', 'Provenance', 'no counted live permission is now refused, refused on appeal or withdrawn', 'PASS' if dead.empty else 'FAIL', len(dead),
      'live register decision and appeal fields', dead[['perm_id', 'ref', 'scheme', 'units', 'status', 'Decision', 'AppealDecision', 'WithdrawnDate']])
pend = x[x.status.eq('In planning') & x.Decision.fillna('').str.upper().isin(['CONDITIONAL', 'UNCONDITIONAL', 'REFUSED'])]
check('P2b', 'Provenance', '"In planning" schemes that have since been decided', 'WARN' if len(pend) else 'PASS', len(pend), 'update status on next rebuild',
      pend[['perm_id', 'ref', 'scheme', 'units', 'Decision', 'AppealDecision']])
from rel import desc_units
x['reg_units'] = pd.to_numeric(x.NumResidentialUnits, errors='coerce')
x['desc_units'] = pd.to_numeric(pd.Series([desc_units(t) if isinstance(t, str) else None for t in x.DevelopmentDescription]), errors='coerce').values
x['best_reg'] = x[['reg_units', 'desc_units']].max(axis=1)
du = x[x.best_reg.gt(0) & x.units.gt(0) & ((x.units - x.best_reg).abs() > np.maximum(5, 0.2 * x.units))]
check('P3', 'Provenance', 'units agree with the planning register (within 20% or 5)', 'PASS' if du.empty else 'WARN', len(du),
      'against the larger of the register unit field and the number in the description', du[['perm_id', 'ref', 'scheme', 'units', 'reg_units', 'desc_units', 'source', 'relationship']])

# P4 appeal cases: SHD / ABP references still granted on the board's case page
ab = json.load(open('abp.json')) if os.path.exists('abp.json') else {}
abr = pp[pp.planning_ref.fillna('').str.startswith('ABPREF')]
rows = []
for _, r in abr.iterrows():
    case = r.planning_ref[6:]
    try:
        t = re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', requests.get(f'https://www.pleanala.ie/en-ie/case/{case}', timeout=60).text))
        dec = re.search(r'Decision (.{0,60}?) Date signed', t)
        hist = 'quash' in t.lower()
        rows.append(dict(perm_id=r.perm_id, case=case, units=r.units, decision=dec.group(1).strip() if dec else '?', quashed=hist))
    except Exception as e:
        rows.append(dict(perm_id=r.perm_id, case=case, units=r.units, decision=f'fetch failed: {e}', quashed=None))
A_ = pd.DataFrame(rows)
badab = A_[~A_.decision.str.contains('Grant', case=False, na=False) | A_.quashed.fillna(False)] if len(A_) else A_
check('P4', 'Provenance', 'appeal-board (SHD) permissions still granted and not quashed', 'PASS' if badab.empty else 'FAIL', len(badab),
      f'{len(A_)} cases checked live on pleanala.ie', A_)

# P5 building control: not-started schemes that have a commencement notice (status lag)
bc = pd.read_pickle('bcms_lim.pkl')
bc['pp'] = bc.CN_Planning_Permission_Number.astype(str).str.replace(r'[^0-9]', '', regex=True)
cn = bc[bc.pp.str.len() >= 4].dropna(subset=['CN_Commencement_Date']).groupby('pp').CN_Commencement_Date.min()
ns = pp[pp.status.isin(['Permitted – not started', 'Stalled', 'Expired – not started'])].copy()
ns['first_commencement'] = ns.ref.fillna('').str.replace(r'[^0-9]', '', regex=True).replace('', None).map(cn)
d = ns[ns.first_commencement.notna()][['perm_id', 'ref', 'scheme', 'units', 'status', 'first_commencement']]
check('P5', 'Provenance', 'not-started schemes with a building control commencement notice', 'WARN' if len(d) else 'PASS', len(d), 'status may be out of date', d)

# P6 hand-entered inputs: each traced to its source
man = []
src = open('build_supply_gpkg.py').read()
for fid, why in re.findall(r"(\d+): '(Reinstated[^']*)'", src):
    man.append(dict(input='Reinstated council audit parcel', key=f'initial audit fid {fid}', value=why))
dsrc = open('delivery_data.py').read()
for ref, why in re.findall(r"'(\d+)': '((?:Granted)[^']*)'", dsrc):
    man.append(dict(input='Scheme added from register (not in CIS)', key=ref, value=why))
for ref in re.findall(r"'(\d{4,8})'", re.search(r"KPMG_ADD = \[([^\]]*)\]", open('sources.py').read()).group(1)):
    man.append(dict(input='Scheme added from KPMG audit', key=ref, value='live in register; missing from CIS exports'))
mx = pd.read_csv('planning_file_mix.csv', dtype={'ref': str}) if os.path.exists('planning_file_mix.csv') else pd.DataFrame()
for _, r in mx.iterrows():
    beds = r[['studio', 'b1', 'b2', 'b3', 'b4', 'b5']].sum()
    ok = beds == r.total and r.houses + r.apartments == r.total
    man.append(dict(input='Unit mix from planning file', key=r.ref, value=f'{r.source}; beds {beds} houses+apts {r.houses + r.apartments} total {r.total}' + ('' if ok else ' – DOES NOT ADD UP')))
M = pd.DataFrame(man)
# verify the added register schemes still match the live register
regadd = M[M.input.str.startswith('Scheme added from register')].key.tolist()
bad = L[L.ApplicationNumber.isin(regadd) & ~L.Decision.isin(['CONDITIONAL', 'UNCONDITIONAL'])]
bad_mix = M[M.value.str.contains('DOES NOT ADD UP')]
check('P6', 'Provenance', 'hand-entered inputs listed and internally consistent', 'PASS' if bad.empty and bad_mix.empty else 'FAIL', len(M),
      f'{len(M)} hand-entered items; {len(bad_mix)} mix rows not adding up; {len(bad)} added schemes no longer granted', M)
mixu = pp[pp.mix_source.eq('Planning file documents')]
d = mixu[(mixu[['mix_studio', 'mix_1bed', 'mix_2bed', 'mix_3bed', 'mix_4bed', 'mix_5plus']].sum(axis=1) - mixu.units).abs() > 0]
check('P7', 'Provenance', 'planning-file unit mixes match counted units', 'PASS' if d.empty else 'WARN', len(d), 'known differences are flagged in planning_file_mix.csv',
      d[['perm_id', 'planning_ref', 'units', 'mix_1bed', 'mix_2bed', 'mix_3bed', 'mix_4bed']])

# ======================================================================== E: external benchmarks (CSO)
def cso(table):
    j = requests.get(f'https://ws.cso.ie/public/api.restful/PxStat.Data.Cube_API.ReadDataset/{table}/JSON-stat/2.0/en', timeout=120).json()
    dims = j['id']
    cats = [list(j['dimension'][d]['category']['index']) if isinstance(j['dimension'][d]['category']['index'], list)
            else sorted(j['dimension'][d]['category']['index'], key=j['dimension'][d]['category']['index'].get) for d in dims]
    labs = [[j['dimension'][d]['category']['label'][c] for c in cs] for d, cs in zip(dims, cats)]
    idx = pd.MultiIndex.from_product(labs, names=dims)
    return pd.Series(j['value'], index=idx, dtype=float).reset_index(name='value')
try:
    n = cso('NDQ06')
    tcol = [c for c in n.columns if c.startswith('TLIST')][0]
    typ = [c for c in n.columns if c.startswith('C02342')][0]
    la = [c for c in n.columns if c.startswith('C03789')][0]
    n = n[n[la].str.contains('Limerick')]
    n['year'] = n[tcol].str[:4].astype(int)
    cy = n.pivot_table(index='year', columns=typ, values='value', aggfunc='sum')
    nq = n[n.year.eq(n.year.max())][tcol].nunique()
    # GeoDirectory new residential addresses by year created (independent count of new homes)
    gd = gpd.read_file('in_geodir/limerick_geod_residential_points.gpkg', columns=['create_dat', 'under_cons', 'building_t'])
    gd['created'] = pd.to_datetime(gd.create_dat, dayfirst=True, errors='coerce')
    gy = gd[gd.under_cons.ne('Y')].groupby(gd.created.dt.year).size()
    # register: CIS-dated completions of counted schemes by year
    sp = gpd.read_file(SUP, layer='permissions')
    sp['cy'] = pd.to_datetime(sp.completion, errors='coerce').dt.year
    ry = sp[sp.status.eq('Complete')].groupby('cy').units.sum()
    E = cy.join(gy.rename('GeoDirectory new addresses'), how='left').join(ry.rename('Register scheme completions (CIS-dated)'), how='left')
    E = E.loc[2019:]
    E['GeoDirectory / CSO all types'] = (E['GeoDirectory new addresses'] / E['All house types']).round(2)
    E['Register / CSO scheme + apartment'] = (E['Register scheme completions (CIS-dated)'] / (E['Scheme house'] + E['Apartment'])).round(2)
    full = E.drop(index=[y for y in E.index if y == n.year.max() and nq < 4], errors='ignore')
    over = full[full['Register / CSO scheme + apartment'] > 1.05]
    gdr = full['GeoDirectory / CSO all types'].median()
    check('E1', 'External', 'register completions never exceed CSO scheme + apartment completions', 'PASS' if over.empty else 'FAIL', len(over),
          'register covers tracked schemes only, so it should sit below CSO', E.reset_index())
    check('E2', 'External', 'GeoDirectory new addresses track CSO completions', 'PASS' if 0.7 <= gdr <= 1.3 else 'WARN', 0,
          f'median ratio {gdr:.2f} over {len(full)} full years (validates GeoDirectory as the completion signal)')
except Exception as e:
    check('E1', 'External', 'CSO completions benchmark', 'WARN', 0, f'CSO fetch failed: {e}')
if os.path.exists('cso_recon.pkl'):
    Y = pd.read_pickle('cso_recon.pkl')['years']
    fy = Y[Y.quarters == 4]
    r = fy['Register total'].sum() / fy['CSO total'].sum()
    check('E3', 'External', 'planning register residential grants vs CSO units permitted (BHQ17)', 'PASS' if 0.85 <= r <= 1.15 else 'WARN', 0,
          f'{fy["Register total"].sum():,.0f} vs {fy["CSO total"].sum():,.0f} over {len(fy)} full years (ratio {r:.2f})', Y.reset_index())

# ======================================================================== C: county audit readiness
za = ps[ps.parcel_source.str.startswith('Zoned land')]
au = ps[ps.parcel_source.eq('Council audit – county')]
if len(au):
    j = gpd.overlay(za[['parcel_id', 'geometry']], au[['parcel_id', 'geometry']], how='intersection', keep_geom_type=True)
    j = j[j.area > 100]
    check('C1', 'County audit', 'county audit loaded; no interim zoned-land parcel overlaps it', 'WARN' if len(j) else 'PASS', len(j),
          f'{len(au)} county audit parcels, {au.capacity_used.sum():,.0f} homes', pd.DataFrame(j.drop(columns='geometry')))
    d = za[za.scope.eq('County')].groupby('plan_name').agg(parcels=('parcel_id', 'size'), ha=('mapped_area_ha', 'sum'), homes=('capacity_used', 'sum')).reset_index()
    check('C2', 'County audit', 'county settlements still on interim zoning estimates (not in the council audit)', 'WARN' if len(d) else 'PASS', int(d.parcels.sum()),
          f'{d.homes.sum():,.0f} homes in {len(d)} plan areas: ask the council for audits of these towns', d)
    d = au[au.notes.fillna('').ne('')][['parcel_id', 'local_name', 'zoning', 'developable_area_ha', 'capacity_used', 'capacity_basis', 'notes']]
    check('C3', 'County audit', 'county audit sites changed on loading (clipped by later consents, or no density given)', 'WARN' if len(d) else 'PASS', len(d),
          'the audit was drawn for the 2022 plan; consents granted since keep their own land', d)
else:
    check('C1', 'County audit', 'county settlements audit loaded', 'WARN', 0, f'not yet received; {len(za)} interim zoned-land parcels ({za.capacity_units.sum():,.0f} homes) stand in for it')

# ======================================================================== QA sample for a person
rng = np.random.default_rng(20261015)
samp = []
for (srcg, st), g_ in pp.assign(srcg=pp.source.fillna('').str.split(' ').str[0]).groupby(['srcg', 'status']):
    k = min(len(g_), max(1, round(30 * len(g_) / len(pp))))
    samp.append(g_.iloc[rng.choice(len(g_), k, replace=False)])
S = pd.concat(samp).merge(L[['ApplicationNumber']].assign(on_register='yes'), left_on='ref', right_on='ApplicationNumber', how='left')
S['register_link'] = np.where(S.ref.fillna('').str.fullmatch(r'\d{4,8}'), 'https://eplanning.ie/LimerickCCC/AppFileRefDetails/' + S.ref.fillna('') + '/0', '')
S = S[['perm_id', 'parcel_id', 'planning_ref', 'scheme', 'units', 'status', 'status_source', 'mix_source', 'catchment', 'source', 'register_link']]
for c in ['units correct (Y/N)', 'status correct (Y/N)', 'location correct (Y/N)', 'checked by', 'note']:
    S[c] = ''
detail['QA sample'] = S

# ======================================================================== write
R = pd.DataFrame(results)
with pd.ExcelWriter(OUT) as w:
    pd.DataFrame({'Validation run': [RUN], 'Register': [REG], 'Checks': [len(R)], 'PASS': [(R.status == 'PASS').sum()],
                  'WARN': [(R.status == 'WARN').sum()], 'FAIL': [(R.status == 'FAIL').sum()]}).to_excel(w, sheet_name='Summary', index=False, startrow=0)
    R.to_excel(w, sheet_name='Summary', index=False, startrow=3)
    for k, v in detail.items():
        v.to_excel(w, sheet_name=k[:31], index=False)
print('\n', R.status.value_counts().to_dict(), '->', OUT)
