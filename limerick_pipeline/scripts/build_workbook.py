"""Build the Limerick residential pipeline workbook (v2).

Counting unit: the application, after primacy rules decide which applications count.
Sites group related applications for review. Bedroom mix filled from register descriptions
where CIS/KPMG mix is blank.

Inputs (same folder): lp.xlsx (source workbook), reg.pkl (pull_register.py), abp.json (fetch_abp_cases.py)
"""
import re, json
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from beds import parse as parse_beds
from rel import classify, desc_units

REPORT_DATE = pd.Timestamp('2026-09-29')
OUT = 'Limerick_Residential_Pipeline_Simplified.xlsx'

x = pd.read_excel('lp.xlsx', sheet_name=None, dtype={'Reference': str})
full = x['Full Project Pipeline'].copy()
from sources import kpmg_records, pbsa_beds, SOURCE as KPMG_SOURCE
_k = kpmg_records(full.columns)
_k['Reference'] = _k['Reference'].astype(str)
full = pd.concat([full, _k], ignore_index=True)
from sources import fix_coords
full = fix_coords(full)
ex = x['Existing Site Pipeline'].copy()
lowconf_sites = set(x['Existing Site Review']['Site Group ID'])
reg = pd.read_pickle('reg.pkl')
reg['ref'] = reg.ApplicationNumber.astype(str).str.strip()
reg = reg.drop_duplicates('ref')
REG = reg.set_index('ref')
abp = json.load(open('abp.json'))

full['ref'] = full.Reference.fillna('').str.strip()
full['Units'] = pd.to_numeric(full.Units, errors='coerce')
for c in ['Application Date', 'Decision Date', 'Start Date', 'Finish Date', 'Last Updated']:
    full[c] = pd.to_datetime(full[c], errors='coerce')


def nz(v):
    try:
        if v is None or pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, str) and v == '':
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating,)):
        return float(v)
    return v


def dist(lat1, lon1, lat2, lon2):
    return np.sqrt(((lat1 - lat2) * 111000) ** 2 + ((lon1 - lon2) * 67000) ** 2)


def outcome(r):
    ad = str(r.AppealDecision or '').upper()
    st = str(r.ApplicationStatus or '').upper()
    dc = str(r.Decision or '').upper()
    if ad in ('CONDITIONAL', 'MODIFIED'):
        return 'Granted on appeal'
    if ad == 'REFUSED':
        return 'Refused on appeal'
    if 'WITHDRAWN' in st:
        return 'Withdrawn'
    if 'INCOMPLETE' in st:
        return 'Invalid'
    if dc == 'REFUSED':
        return 'Refused'
    if dc in ('CONDITIONAL', 'UNCONDITIONAL'):
        return 'Granted'
    if st.startswith('APPEALED') or st == 'LEAVE TO APPEAL':
        return 'Under appeal'
    return 'Pending'


# extension of duration lookup by parent file number quoted in description
eod = {}
for _, r in reg[reg.ApplicationType == 'EXTENSION OF DURATION'].sort_values('ReceivedDate').iterrows():
    for yy, n in re.findall(r'(\d{2})\s*/\s*(\d+)', str(r.DevelopmentDescription)):
        eod[yy + n] = r

CISREFS = set(full.ref)
regpts = reg[reg.lat.notna() & (reg.NumResidentialUnits.fillna(0) >= 3) & (reg.ApplicationType != 'EXTENSION OF DURATION')]


def later_nearby(lat, lon, after, own):
    if lat is None or pd.isna(lat) or pd.isna(after):
        return ''
    d = dist(regpts.lat, regpts.lon, lat, lon)
    c = regpts[(d < 150) & (regpts.ReceivedDate > after) & (regpts.ref != own)]
    if c.empty:
        return ''
    r = c.sort_values(['NumResidentialUnits', 'ReceivedDate']).iloc[-1]
    return f"{r.ref} ({r.ReceivedDate:%d/%m/%Y}, {outcome(r)}, {int(r.NumResidentialUnits)} units, {'in CIS' if r.ref in CISREFS else 'NOT in CIS'})"


# ---------------- register match per application ----------------
recs = []
for _, a in full.iterrows():
    rec = dict(found='No reference', reg_ref='', status='', decision='', appeal='', outcome='No reference',
               reg_dec=None, reg_exp=None, eod='', exp_final=None, reg_units=None, reg_desc='', desc_full=None,
               regtype='', locdiff=None, later='', link='')
    ref = a.ref
    if ref and ref in REG.index:
        r = REG.loc[ref]
        o = outcome(r)
        rec.update(found='Council register', reg_ref=ref, status=r.ApplicationStatus, decision=r.Decision,
                   appeal=r.AppealDecision or '', outcome=o, reg_dec=r.DecisionDate, reg_exp=r.ExpiryDate,
                   reg_units=(r.NumResidentialUnits if pd.notna(r.NumResidentialUnits) and r.NumResidentialUnits > 0 else None),
                   reg_desc=str(r.DevelopmentDescription)[:300], desc_full=r.DevelopmentDescription,
                   regtype=r.ApplicationType, link=r.LinkAppDetails or '')
        if pd.notna(a.Latitude) and pd.notna(r.lat):
            rec['locdiff'] = round(float(dist(a.Latitude, a.Longitude, r.lat, r.lon)))
        rec['exp_final'] = r.ExpiryDate
        if ref in eod:
            e = eod[ref]
            rec['eod'] = f"{e.ref}: {outcome(e)}" + (f", expires {e.ExpiryDate:%d/%m/%Y}" if pd.notna(e.ExpiryDate) else '')
            if outcome(e) in ('Granted', 'Granted on appeal') and pd.notna(e.ExpiryDate):
                rec['exp_final'] = e.ExpiryDate
        if o in ('Refused', 'Refused on appeal', 'Withdrawn', 'Invalid', 'Quashed'):
            rec['later'] = later_nearby(a.Latitude, a.Longitude, r.ReceivedDate, ref)
    elif ref.startswith('ABPREF') and ref[6:] in abp:
        b = abp[ref[6:]]
        signed = pd.to_datetime(b['signed'], dayfirst=True)
        yrs = 10 if 'ten year' in b['desc'].lower() else 5
        o = 'Quashed' if 'quash' in b.get('history', '').lower() else ('Granted' if 'grant' in b['decision'].lower() else ('Refused' if 'refus' in b['decision'].lower() else 'Pending'))
        m = re.search(r'(\d+)\s*no\.?\s*(?:residential units|Build to Rent apartments)', b['desc'])
        rec.update(found='An Coimisiún Pleanála', reg_ref='ABP-' + ref[6:], decision=b['decision'], outcome=o,
                   reg_dec=signed, reg_exp=signed + pd.DateOffset(years=yrs) - pd.Timedelta(days=1),
                   exp_final=signed + pd.DateOffset(years=yrs) - pd.Timedelta(days=1),
                   reg_units=int(m.group(1)) if m else None, reg_desc=b['desc'][:300], desc_full=b['desc'],
                   link=b['url'], status=f'{yrs}-year permission (expiry = date signed + {yrs} years)')
    elif ref:
        rec.update(found='Not in council register', outcome='Not in council register')
    recs.append(rec)
A = pd.concat([full.reset_index(drop=True), pd.DataFrame(recs)], axis=1)
A['desc'] = [d if isinstance(d, str) else h for d, h in zip(A.desc_full, A['Project Heading'])]
A['dunits'] = [desc_units(d) if isinstance(d, str) else None for d in A.desc_full]

# ---------------- sites ----------------
pid2site = {}
for _, s in ex.iterrows():
    for p in str(s['Related Project IDs']).split(','):
        pid2site[int(p)] = s['Site Group ID']
granted = A[A['Source Dataset'].str.startswith('Granted')]
up = A[~A['Source Dataset'].str.startswith('Granted')]

SETTLE = ['Mungret', 'Annacotty', 'Patrickswell', 'Newcastle West', 'Kilmallock', 'Adare', 'Bruff', 'Cappamore',
          'Rathkeale', 'Castleconnel', 'Caherconlish', 'Abbeyfeale']
CITY = ['Limerick City', 'Moyross', 'Thomondgate', 'Thomondghate', 'Singland', 'Dooradoyle', 'Raheen', 'Southill',
        'Castletroy', 'Corbally', 'Caherdavin', 'Crossagalla', 'Kennedy Park', "King's Island", 'Dock Road',
        'South Circular', 'Mary Street', 'Parnell', 'New Road', 'Dublin Road', 'Rosbrien', 'Limerick']


def settlement(addr):
    for s in SETTLE:
        if s.lower() in addr.lower():
            return 'Castleconnell' if s == 'Castleconnel' else s
    for c in CITY:
        if c.lower() in addr.lower():
            return 'Limerick City and Suburbs'
    return 'Other County Limerick'


def addr(r):
    parts = [r[f'Project Site Address Line {i}'] for i in range(1, 6)]
    return ', '.join(str(p).strip() for p in parts if pd.notna(p) and str(p).strip())


sites = []
for _, s in ex.iterrows():
    sites.append(dict(sid=s['Site Group ID'], origin='Granted/delivery export (grouped)', heading=s['Site / Project Heading'],
                      address=s['Address'], settlement=s['Settlement'], lat=s['Latitude'], lon=s['Longitude'],
                      stage=s['Current Stage'], done=s['Units Completed Stated'], overlap='',
                      lowconf='Yes' if s['Site Group ID'] in lowconf_sites else 'No'))
seen = []
for _, r in up.sort_values('Application Date').iterrows():
    if r['Residential Pipeline Included'] != 'Yes':
        continue
    sid = f"U-{r['Project Id']}"
    pid2site[r['Project Id']] = sid
    notes = []
    if pd.notna(r.Latitude):
        d = dist(granted.Latitude, granted.Longitude, r.Latitude, r.Longitude)
        for j in d[d < 150].sort_values().index[:2]:
            g = granted.loc[j]
            notes.append(f"{pid2site[g['Project Id']]} ({int(d[j])}m): {g['Project Heading'][:45]}, {g['Pipeline Stage']}, "
                         f"{int(g.Units) if pd.notna(g.Units) else '?'} units")
        for o in seen:
            dd = dist(o[1], o[2], r.Latitude, r.Longitude)
            if dd < 150:
                notes.append(f"{o[0]} ({int(dd)}m): {o[3][:45]}, {o[4]}, {o[5]} units")
        seen.append((sid, r.Latitude, r.Longitude, r['Project Heading'], r['Pipeline Stage'],
                     int(r.Units) if pd.notna(r.Units) else '?'))
    ad = addr(r)
    sites.append(dict(sid=sid, origin='Latest upstream export', heading=r['Project Heading'], address=ad,
                      settlement=settlement(ad), lat=r.Latitude, lon=r.Longitude, stage=r['Pipeline Stage'],
                      done=None, overlap='; '.join(notes), lowconf='No'))
S = pd.DataFrame(sites)
A['site'] = A['Project Id'].map(pid2site).fillna('')
A['settlement'] = A.site.map(S.set_index('sid').settlement).fillna('')

# units completed (stated) sit on the site; assign to the largest Part Complete application
A['done'] = None
for s in S.itertuples():
    if pd.notna(s.done) and s.done > 0:
        pc = A[(A.site == s.sid) & (A['Pipeline Stage'] == 'Part Complete')]
        if len(pc):
            A.loc[pc.Units.idxmax(), 'done'] = s.done

# ---------------- relationships ----------------
D = pd.DataFrame(dict(ref=A.ref.where(A.ref != '', 'PID' + A['Project Id'].astype(str)), site=A.site,
                      date=A['Application Date'], units=A.Units, dunits=A.dunits, outcome=A.outcome,
                      stage=A['Pipeline Stage'], desc=A.desc, lat=A.Latitude, lon=A.Longitude, regtype=A.regtype,
                      included=A['Residential Pipeline Included']))
rel = classify(D)
for k, n in enumerate(['rel', 'rel_ref', 'rel_basis', 'rel_conf']):
    A[n] = [rel[i][k] for i in A.index]
A['rel_ref'] = A.rel_ref.str.replace('^PID', 'Project ', regex=True)

# CIS sometimes holds one Part 8 twice: under its own project number and under the council file number (yy/8nnn)
_lab = A[A.ref.str.contains(r'part\s*8|PT8', case=False, na=False) & (A.found != 'Council register') & A.rel.isin(['Primary', 'Additional application on site', 'Additional phase'])]
_r8 = A[A.ref.str.fullmatch(r'\d{2}8\d{3}', na=False) & ~A.rel.isin(['Not live', 'Superseded', 'Excluded', 'Duplicate'])]
for i, a in _lab.iterrows():
    if pd.isna(a.Latitude) or pd.isna(a.Units):
        continue
    d = np.sqrt(((_r8.Latitude - a.Latitude) * 111000) ** 2 + ((_r8.Longitude - a.Longitude) * 67000) ** 2)
    ok = (d < 100) & ((_r8.Units - a.Units).abs() <= 0.1 * a.Units)
    if ok.any():
        j = d[ok].idxmin()
        A.loc[i, ['rel', 'rel_ref', 'rel_basis', 'rel_conf']] = ['Duplicate', _r8.loc[j, 'ref'],
            f"Same Part 8 recorded by CIS under its project number and council file {_r8.loc[j, 'ref']} ({int(d[j])}m, same units)", 'High']

# bedspaces for granted student schemes from the PBSA file
_beds = pbsa_beds()
A['Student Bedspaces'] = [_beds.get(r, b) for r, b in zip(A.ref, A['Student Bedspaces'])]


# ---------------- dwelling type and bedroom mix ----------------
def dtype(r):
    h = str(r['Project Heading']).lower()
    if 'student' in h or (pd.notna(r['Student Bedspaces']) and r['Student Bedspaces'] > 0):
        return 'Student accommodation'
    return 'General'


A['dwtype'] = A.apply(dtype, axis=1)
BEDCOLS = ['Studio', '1 Bed', '2 Bed', '3 Bed', '4 Bed', '5+ Bed']
mix_rows = []
for _, r in A.iterrows():
    m = dict(src='Not available', note='', houses=None, apts=None, unk=None, **{b: None for b in BEDCOLS})
    cisbeds = [r[f'{b} Bed Apartments'] for b in ('1', '2', '3', '4', '5+')] + [r[f'{b} Bed Houses'] for b in ('1', '2', '3', '4', '5+')]
    if any(pd.notna(v) for v in cisbeds):
        for k, b in enumerate(['1 Bed', '2 Bed', '3 Bed', '4 Bed', '5+ Bed']):
            ap = r[f'{("5+" if b == "5+ Bed" else b[0])} Bed Apartments']
            ho = r[f'{("5+" if b == "5+ Bed" else b[0])} Bed Houses']
            tot = (0 if pd.isna(ap) else ap) + (0 if pd.isna(ho) else ho)
            m[b] = tot if tot else None
        m['houses'] = nz(r['Total Houses'])
        m['apts'] = nz(r['Total Apartments'])
        m['src'] = 'CIS / KPMG'
    elif isinstance(r.desc_full, str):
        mix, tot = parse_beds(r.desc_full)
        if tot:
            targets = {t for t in (r.dunits, r.Units, r.reg_units) if pd.notna(t) and t}
            if tot in targets:
                for (k, b), n in mix.items():
                    col = BEDCOLS[min(b, 5)]
                    m[col] = (m[col] or 0) + n
                m['houses'] = sum(n for (k, b), n in mix.items() if k == 'house') or None
                m['apts'] = sum(n for (k, b), n in mix.items() if k == 'apt') or None
                m['unk'] = sum(n for (k, b), n in mix.items() if k == 'unit') or None
                m['src'] = 'Register description'
            else:
                m['note'] = (f"Partial parse not used ({tot} units found vs {int(max(targets)) if targets else '?'} stated): "
                             + ', '.join(f"{n}x{b if b else 'studio'}-bed {k}" for (k, b), n in sorted(mix.items(), key=lambda z: z[0][1])))
                m['src'] = 'Not available (register text incomplete)'
        else:
            m['src'] = 'Not available (no bedroom detail in register text)'
    else:
        m['src'] = 'Not available (not in register)'
    mix_rows.append(m)
M = pd.DataFrame(mix_rows)
A = pd.concat([A, M], axis=1)

import os
SHF = {}
if os.path.exists('shcp_feed.pkl'):
    _f = pd.read_pickle('shcp_feed.pkl')
    SHF = {int(k): {kk: (None if (isinstance(vv, float) and np.isnan(vv)) else vv) for kk, vv in v.items()}
           for k, v in _f.set_index('pid').to_dict('index').items() if str(k).isdigit()}

# =============================== WORKBOOK ===============================
wb = Workbook()
F = 'Arial'
HFILL = PatternFill('solid', fgColor='1F3A5F')
GFILL = PatternFill('solid', fgColor='375623')
PFILL = PatternFill('solid', fgColor='5B3A8C')
BFILL = PatternFill('solid', fgColor='7F6000')
INFILL = PatternFill('solid', fgColor='FFF2CC')
RED = PatternFill('solid', fgColor='F8CBAD')
AMB = PatternFill('solid', fgColor='FFE699')
DATE = 'dd/mm/yyyy'
NUM = '#,##0;(#,##0);-'


def header(ws, row, cols, fill=HFILL):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(name=F, bold=True, color='FFFFFF', size=10)
        cell.fill = fill if not isinstance(fill, dict) else fill.get(c, HFILL)
        cell.alignment = Alignment(wrap_text=True, vertical='top')


def inp(cell):
    cell.fill = INFILL
    cell.font = Font(name=F, size=10, color='0000FF')


# ---------------- Method ----------------
Mw = wb.active
Mw.title = 'Method'
Mw['A1'] = 'Limerick Residential Pipeline – Method and Assumptions'
Mw['A1'].font = Font(name=F, bold=True, size=14)
Mw['A3'] = 'Assumptions (edit the yellow cells)'
Mw['A3'].font = Font(name=F, bold=True, size=11)
assump = [
    ('Report date', REPORT_DATE, 'Date the register was queried. Used for staleness and expiry tests.'),
    ('Stale threshold (months)', 24, 'A counted Permission Granted / Commencement Expected application with no CIS update for longer than this is reported as Constrained / uncertain until checked.'),
    ('Overlap search radius (m)', 150, 'Used at build time to flag upstream records near another record. Changing it does not re-run the search.'),
    ('Location mismatch threshold (m)', 500, 'Register point more than this far from the CIS point is flagged as a possible reference mismatch.'),
]
for i, (k, v, n) in enumerate(assump, start=4):
    Mw.cell(row=i, column=1, value=k).font = Font(name=F, bold=True, size=10)
    c = Mw.cell(row=i, column=2, value=nz(v))
    inp(c)
    if isinstance(v, pd.Timestamp):
        c.number_format = DATE
    t = Mw.cell(row=i, column=3, value=n)
    t.font = Font(name=F, size=10)
    t.alignment = Alignment(wrap_text=True, vertical='top')
RD, STALE, LOCM = 'Method!$B$4', 'Method!$B$5', 'Method!$B$7'

notes = [
    ('Sources', ''),
    ('CIS granted/delivery export', 'Granted/delivery records and their site grouping from the Existing Site Pipeline tab of Limerick_Residential_Pipeline_Full_Build.xlsx (226 records in 153 sites).'),
    ('CIS upstream export', 'Export 20260923095518.xls as de-duplicated in the source workbook (55 records; 3 non-residential frameworks excluded).'),
    ('Planning register', 'National Planning Applications dataset (Department of Housing, Local Government and Heritage), ArcGIS feature service IrishPlanningApplications, Limerick County Council records (12,835, received 2017 onwards). Queried 29/09/2026.'),
    ('An Coimisiún Pleanála', 'SHD case pages at pleanala.ie/en-ie/case/<number>, read 29/09/2026 for the 6 records with ABP references.'),
    ('', ''),
    ('Counting (primacy)', ''),
    ('Counting unit', 'The application. Each application gets a Relationship; only Primary, Additional phase, Additional application on site, Replaces earlier and Extension of duration (parent not in CIS) count. Sites group applications for review; site totals are sums of counted applications.'),
    ('Extension of duration', 'Register EoD applications, or descriptions starting "extension of permission/duration", link to the parent file number in the description (or to the most similar earlier application on the site). Units stay on the parent; the parent expiry is extended.'),
    ('Replaces earlier', 'A later application that (a) modifies a quoted earlier permission and restates at least half its units, (b) is consequent on a quoted outline permission, or (c) repeats an earlier live application nearby (description text match >= 85% with units within 25%, or units within 10% with text match >= 50%). The earlier application becomes Superseded.'),
    ('Amendment within parent', 'Modifies a quoted earlier permission but covers less than half its units. Its units are treated as already counted in the parent.'),
    ('Additional phase', 'Quotes an earlier permission without modification wording (e.g. phase 2, connecting to services of), or modifies a completed scheme or one with no unit count. Counted.'),
    ('Additional application on site', 'Other live applications exist on the same site or within 100m, but no link could be found in the descriptions. Counted by default and marked Low confidence – these are the main remaining double-count risk and are in the Review Queue.'),
    ('Not live', 'Refused, refused on appeal, withdrawn or invalid in the register. Not counted.'),
    ('Analyst override', 'Set Analyst Relationship on the Applications tab to override the automatic relationship. Final Relationship and all totals update.'),
    ('', ''),
    ('Reporting tier and units', ''),
    ('Reporting Tier', 'Per counted application. Analyst Status overrides. Otherwise Complete = Delivered; Permission Granted or Commencement Expected that is past expiry or stale = Constrained / uncertain; else the CIS tier.'),
    ('Units', 'CIS units. Remaining = units less stated completions (completions recorded on the site are assigned to its largest Part Complete application). Delivered = units on Complete applications plus stated completions.'),
    ('Student accommodation', 'Student units and bedspaces reported separately and never added to general dwellings.'),
    ('Expiry', 'Register expiry date, replaced by a granted Extension of Duration expiry where one quotes the parent file number; ABP expiry = date signed + 5 (or 10) years; CIS decision date + 5 years as fallback.'),
    ('', ''),
    ('Bedroom mix', ''),
    ('Priority', 'CIS / KPMG bedroom mix where present. Otherwise the register development description is parsed for "N no. X-bed ..." patterns.'),
    ('Acceptance rule', 'A parsed mix is used only when the bedrooms found add up exactly to the stated unit total (description header, CIS units or register units). Partial parses are shown in Mix Note but not used.'),
    ('Coverage', 'Mix Coverage on the Summary is the share of counted units with a known bedroom split. Descriptions that only say "N houses" carry no bedroom data; those need the planning drawings or schedule of accommodation.'),
    ('', ''),
    ('Limitations', ''),
    ('Part 8 and pre-2017', 'Part 8 references and applications before 2017 are not in the national dataset and stay unverified.'),
    ('Register unit counts', 'NumResidentialUnits is blank or zero on many register records.'),
    ('Automatic relationships', 'Rules read the text of descriptions. They cannot see drawings, so phase boundaries and repeat applications on different parts of a large landholding can be misread. Check Low and Medium confidence rows.'),
    ('Legal position', 'Screening only. It does not confirm the legal status of any permission.'),
    ('Social housing report', 'Social Housing Construction Projects Status Report Q1 2026 (Department of Housing, Limerick pages). The report has no planning references or coordinates; only schemes matched with confidence on name, address and units (and consistent dates) are shown in the Social Housing columns. Unmatched and uncertain rows are omitted.'),
    ('Contacts', 'CIS project team names and emails are not carried into this file.'),
]
r0 = 10
for i, (k, v) in enumerate(notes):
    a = Mw.cell(row=r0 + i, column=1, value=k or None)
    b = Mw.cell(row=r0 + i, column=2, value=v or None)
    a.font = Font(name=F, bold=True, size=11 if (k and not v) else 10)
    b.font = Font(name=F, size=10)
    Mw.merge_cells(start_row=r0 + i, start_column=2, end_row=r0 + i, end_column=3)
    b.alignment = Alignment(wrap_text=True, vertical='top')
Mw.column_dimensions['A'].width = 30
Mw.column_dimensions['B'].width = 16
Mw.column_dimensions['C'].width = 110

# ---------------- Applications ----------------
AW = wb.create_sheet('Applications')
groups = [
    (HFILL, ['Project Id', 'Site ID', 'CIS Reference', 'Project Heading', 'Address', 'Settlement', 'CIS Stage',
             'Planning Stage', 'Contract Stage', 'Application Date', 'CIS Decision Date', 'Last Updated',
             'Source Dataset', 'Included in Pipeline', 'Exclusion Reason', 'Dwelling Type', 'Units',
             'Units in Description', 'Student Bedspaces']),
    (PFILL, ['Auto Relationship', 'Related Ref', 'Relationship Basis', 'Confidence', 'Analyst Relationship',
             'Final Relationship', 'Counts?']),
    (HFILL, ['Units Completed (stated)', 'Months Since Update', 'Stale?', 'Expiry Used', 'Past Expiry?',
             'Analyst Status', 'Reporting Tier', 'Tier Reason', 'Counted Units', 'Remaining Units', 'Delivered Units',
             'Pipeline General Dwellings', 'Pipeline Student Units', 'Pipeline Student Bedspaces']),
    (BFILL, ['Mix Source', 'Studio', '1 Bed', '2 Bed', '3 Bed', '4 Bed', '5+ Bed', 'Mix Total', 'Houses',
             'Apartments', 'Type Unknown', 'Mix Note']),
    (GFILL, ['Register Source', 'Register Ref', 'Register Status', 'Register Decision', 'Appeal Decision',
             'Register Outcome', 'Register Decision Date', 'Register Expiry', 'Extension of Duration',
             'Expiry incl. EoD', 'Register Units', 'Location Diff (m)', 'Later Application Nearby',
             'Register Description', 'Register Link', 'Decision Date Diff (days)', 'Units Check', 'Check Result']),
    (PatternFill('solid', fgColor='843C0C'), ['Social Housing Project No.', 'Social Housing Programme', 'Social Housing AHB',
             'Social Housing Stage', 'Social Housing Stage Quarter', 'Public Delivery']),
]
acols, afill = [], {}
for fill, cols in groups:
    acols += cols
    for c in cols:
        afill[c] = fill
CA = {n: L(i) for i, n in enumerate(acols, 1)}
header(AW, 1, acols, afill)
N = len(A)
AL = N + 1
COUNTED = ['Primary', 'Additional phase', 'Additional application on site', 'Replaces earlier',
           'Extension of duration (parent not in CIS)']
RELS = COUNTED + ['Superseded', 'Amendment within parent', 'Extension of duration', 'Not live', 'Duplicate', 'Excluded']
order = A.assign(_s=A.site.replace('', 'ZZZ')).sort_values(['_s', 'Application Date'], na_position='last').index
for i, idx in enumerate(order, start=2):
    r = A.loc[idx]
    c = lambda k: f'{CA[k]}{i}'
    vals = {
        'Project Id': r['Project Id'], 'Site ID': r.site, 'CIS Reference': r.Reference, 'Project Heading': r['Project Heading'],
        'Address': addr(r), 'Settlement': r.settlement, 'CIS Stage': r['Pipeline Stage'], 'Planning Stage': r['Planning Stage'],
        'Contract Stage': r['Contract Stage'], 'Application Date': r['Application Date'], 'CIS Decision Date': r['Decision Date'],
        'Last Updated': r['Last Updated'], 'Source Dataset': r['Source Dataset'], 'Included in Pipeline': r['Residential Pipeline Included'],
        'Exclusion Reason': r['Exclusion Reason'], 'Dwelling Type': r.dwtype, 'Units': r.Units, 'Units in Description': r.dunits,
        'Student Bedspaces': r['Student Bedspaces'], 'Auto Relationship': r.rel, 'Related Ref': r.rel_ref,
        'Relationship Basis': r.rel_basis, 'Confidence': r.rel_conf, 'Units Completed (stated)': r.done,
        'Mix Source': r.src, 'Houses': r.houses, 'Apartments': r.apts, 'Type Unknown': r.unk, 'Mix Note': r.note,
        'Social Housing Project No.': SHF.get(r['Project Id'], {}).get('no'), 'Social Housing Programme': SHF.get(r['Project Id'], {}).get('programme'),
        'Social Housing AHB': SHF.get(r['Project Id'], {}).get('ahb'), 'Social Housing Stage': SHF.get(r['Project Id'], {}).get('stage'),
        'Social Housing Stage Quarter': SHF.get(r['Project Id'], {}).get('stage_quarter'),
        'Public Delivery': (SHF[r['Project Id']]['sh_owner'].replace('Public – ', '').capitalize() + ' (' + SHF[r['Project Id']]['sh_mode'].lower() + ')') if r['Project Id'] in SHF else None,
        'Register Source': r.found, 'Register Ref': r.reg_ref, 'Register Status': r.status, 'Register Decision': r.decision,
        'Appeal Decision': r.appeal, 'Register Outcome': r.outcome, 'Register Decision Date': r.reg_dec,
        'Register Expiry': r.reg_exp, 'Extension of Duration': r.eod, 'Expiry incl. EoD': r.exp_final,
        'Register Units': r.reg_units, 'Location Diff (m)': r.locdiff, 'Later Application Nearby': r.later,
        'Register Description': r.reg_desc, 'Register Link': r.link,
    }
    for b in BEDCOLS:
        vals[b] = r[b]
    for k, v in vals.items():
        AW[c(k)] = nz(v)
    inp(AW[c('Analyst Relationship')])
    inp(AW[c('Analyst Status')])
    st, fr, tier = c('CIS Stage'), c('Final Relationship'), c('Reporting Tier')
    AW[fr] = f'=IF({c("Analyst Relationship")}<>"",{c("Analyst Relationship")},{c("Auto Relationship")})'
    AW[c('Counts?')] = '=IF(OR(' + ','.join(f'{fr}="{x}"' for x in COUNTED) + '),"Yes","No")'
    AW[c('Months Since Update')] = f'=IF({c("Last Updated")}="","",ROUND(({RD}-{c("Last Updated")})/30.44,0))'
    AW[c('Stale?')] = f'=IF({c("Months Since Update")}="","No",IF({c("Months Since Update")}>{STALE},"Yes","No"))'
    AW[c('Expiry Used')] = f'=IF({c("Expiry incl. EoD")}<>"",{c("Expiry incl. EoD")},IF({c("CIS Decision Date")}<>"",EDATE({c("CIS Decision Date")},60),""))'
    pgce = f'OR({st}="Permission Granted",{st}="Commencement Expected",{st}="Stalled or Suspended")'
    AW[c('Past Expiry?')] = f'=IF(AND({c("Expiry Used")}<>"",{pgce}),IF({c("Expiry Used")}<{RD},"Yes","No"),"No")'
    dec = c('Analyst Status')
    # CIS tier lookup from stage
    cistier = (f'IF({st}="On Site","Committed",IF({st}="Part Complete","Committed",IF({st}="Contract Awarded","Committed Pre-Construction",'
               f'IF({st}="Tender","Advanced Pre-Construction",IF({st}="Commencement Expected","Likely",IF({st}="Permission Granted","Potential",'
               f'IF({st}="Stalled or Suspended","Constrained / uncertain",IF({st}="Complete","Delivered","Application Pipeline"))))))))')
    pg2 = f'OR({st}="Permission Granted",{st}="Commencement Expected")'
    AW[tier] = (f'=IF({c("Counts?")}="No","Not counted",IF({dec}="Lapsed or dead","Excluded",IF({dec}="Constrained","Constrained / uncertain",'
                f'IF({dec}="Confirmed live",{cistier},IF({st}="Complete","Delivered",'
                f'IF(AND({pg2},OR({c("Stale?")}="Yes",{c("Past Expiry?")}="Yes")),"Constrained / uncertain",{cistier}))))))')
    AW[c('Tier Reason')] = (f'=IF({c("Counts?")}="No","Relationship: "&{fr},IF({dec}<>"","Analyst: "&{dec},IF({st}="Complete","Complete",'
                            f'IF(AND({pg2},{c("Past Expiry?")}="Yes"),"Past expiry",IF(AND({pg2},{c("Stale?")}="Yes"),"No CIS update in "&{c("Months Since Update")}&" months","As CIS")))))')
    u = c('Units')
    AW[c('Counted Units')] = f'=IF({c("Counts?")}="Yes",N({u}),0)'
    AW[c('Remaining Units')] = f'=IF(OR({c("Counts?")}="No",{st}="Complete"),0,MAX(N({u})-N({c("Units Completed (stated)")}),0))'
    AW[c('Delivered Units')] = f'=IF({c("Counts?")}="No",0,IF({st}="Complete",N({u}),MIN(N({c("Units Completed (stated)")}),N({u}))))'
    notpipe = f'OR({tier}="Not counted",{tier}="Excluded",{tier}="Delivered")'
    dt = c('Dwelling Type')
    AW[c('Pipeline General Dwellings')] = f'=IF(OR({notpipe},{dt}="Student accommodation"),0,{c("Remaining Units")})'
    AW[c('Pipeline Student Units')] = f'=IF(OR({notpipe},{dt}<>"Student accommodation"),0,{c("Remaining Units")})'
    AW[c('Pipeline Student Bedspaces')] = f'=IF({notpipe},0,N({c("Student Bedspaces")}))'
    AW[c('Mix Total')] = '=' + '+'.join(f'N({c(b)})' for b in BEDCOLS)
    AW[c('Decision Date Diff (days)')] = f'=IF(OR({c("CIS Decision Date")}="",{c("Register Decision Date")}=""),"",ROUND({c("Register Decision Date")}-{c("CIS Decision Date")},0))'
    AW[c('Units Check')] = f'=IF(OR(N({c("Register Units")})=0,{u}=""),"Not recorded",IF({c("Register Units")}={u},"Match","Differs"))'
    o = c('Register Outcome')
    built = f'OR({st}="Complete",{st}="On Site",{st}="Part Complete")'
    dead = f'OR({o}="Refused",{o}="Refused on appeal",{o}="Withdrawn",{o}="Invalid",{o}="Quashed")'
    AW[c('Check Result')] = (
        f'=IF({c("Included in Pipeline")}="No","Excluded (non-residential)",IF({o}="No reference","Unverified – no reference",'
        f'IF({o}="Not in council register","Unverified – not in register (Part 8 / other)",'
        f'IF({dead},IF({built},"Conflict – register not live but CIS says built","Conflict – not live in register"),'
        f'IF(AND({st}="Plans Submitted",OR({o}="Granted",{o}="Granted on appeal")),"Stage out of date – now granted",'
        f'IF(AND({st}<>"Plans Submitted",OR({o}="Pending",{o}="Under appeal")),"Check – register still pending",'
        f'IF(N({c("Location Diff (m)")})>{LOCM},"Check – location differs",'
        f'IF(AND({c("Decision Date Diff (days)")}<>"",ABS(N({c("Decision Date Diff (days)")}))>60),"Consistent – decision date differs","Consistent"))))))))')
    if r.link:
        AW[c('Register Link')].hyperlink = r.link
        AW[c('Register Link')].font = Font(name=F, size=10, color='0563C1', underline='single')
dvr = DataValidation(type='list', formula1='"' + ','.join(RELS) + '"', allow_blank=True)
dvs = DataValidation(type='list', formula1='"Confirmed live,Constrained,Lapsed or dead"', allow_blank=True)
AW.add_data_validation(dvr)
AW.add_data_validation(dvs)
dvr.add(f"{CA['Analyst Relationship']}2:{CA['Analyst Relationship']}{AL}")
dvs.add(f"{CA['Analyst Status']}2:{CA['Analyst Status']}{AL}")
for k in ('Application Date', 'CIS Decision Date', 'Last Updated', 'Expiry Used', 'Register Decision Date', 'Register Expiry', 'Expiry incl. EoD'):
    for rr in range(2, AL + 1):
        AW[f'{CA[k]}{rr}'].number_format = DATE
for k in ('Units', 'Units in Description', 'Student Bedspaces', 'Counted Units', 'Remaining Units', 'Delivered Units',
          'Pipeline General Dwellings', 'Pipeline Student Units', 'Pipeline Student Bedspaces', 'Mix Total', 'Houses',
          'Apartments', 'Type Unknown', *BEDCOLS):
    for rr in range(2, AL + 1):
        AW[f'{CA[k]}{rr}'].number_format = NUM
aw = {'Project Heading': 40, 'Address': 30, 'Relationship Basis': 44, 'Auto Relationship': 22, 'Final Relationship': 22,
      'Analyst Relationship': 20, 'Register Description': 50, 'Later Application Nearby': 36, 'Check Result': 34,
      'Extension of Duration': 28, 'Register Link': 28, 'Mix Note': 44, 'Mix Source': 24, 'Tier Reason': 22,
      'Reporting Tier': 20, 'Settlement': 18, 'Source Dataset': 18, 'Exclusion Reason': 20, 'Register Status': 18,
      'Register Outcome': 16, 'Register Source': 16, 'CIS Stage': 16, 'Analyst Status': 14, 'Dwelling Type': 14}
for k, col in CA.items():
    AW.column_dimensions[col].width = aw.get(k, 11)
AW.freeze_panes = 'E2'
AW.auto_filter.ref = f'A1:{L(len(acols))}{AL}'
AW.row_dimensions[1].height = 45
cr, cf = CA['Check Result'], CA['Confidence']
AW.conditional_formatting.add(f'{cr}2:{cr}{AL}', FormulaRule(formula=[f'LEFT({cr}2,8)="Conflict"'], fill=RED))
AW.conditional_formatting.add(f'{cr}2:{cr}{AL}', FormulaRule(formula=[f'OR(LEFT({cr}2,5)="Check",LEFT({cr}2,5)="Stage")'], fill=AMB))
AW.conditional_formatting.add(f'{cf}2:{cf}{AL}', FormulaRule(formula=[f'AND({cf}2="Low",{CA["Analyst Relationship"]}2="")'], fill=AMB))

# ---------------- Sites (roll-up) ----------------
SW = wb.create_sheet('Sites')
scols = ['Site ID', 'Origin', 'Site / Project Heading', 'Address', 'Settlement', 'Latitude', 'Longitude',
         'Most Advanced CIS Stage', 'Applications', 'Counted Applications', 'Counted Units', 'Pipeline General Dwellings',
         'Pipeline Student Units', 'Delivered Units', 'Not-Counted Applications', 'Low-Confidence Links (unreviewed)',
         'Stale / Expired (counted)', 'Possible Overlap', 'Low-Confidence Grouping', 'Review Required',
         'Site Reviewed', 'Site Note']
CS = {n: L(i) for i, n in enumerate(scols, 1)}
header(SW, 1, scols)
SL = len(S) + 1
ar = lambda k: f"Applications!${CA[k]}$2:${CA[k]}${AL}"
for i, s in enumerate(S.itertuples(), start=2):
    c = lambda k: f'{CS[k]}{i}'
    for k, v in [('Site ID', s.sid), ('Origin', s.origin), ('Site / Project Heading', s.heading), ('Address', s.address),
                 ('Settlement', s.settlement), ('Latitude', s.lat), ('Longitude', s.lon), ('Most Advanced CIS Stage', s.stage),
                 ('Possible Overlap', s.overlap), ('Low-Confidence Grouping', s.lowconf)]:
        SW[c(k)] = nz(v)
    sid = c('Site ID')
    SW[c('Applications')] = f'=COUNTIF({ar("Site ID")},{sid})'
    SW[c('Counted Applications')] = f'=COUNTIFS({ar("Site ID")},{sid},{ar("Counts?")},"Yes")'
    SW[c('Counted Units')] = f'=SUMIFS({ar("Counted Units")},{ar("Site ID")},{sid})'
    SW[c('Pipeline General Dwellings')] = f'=SUMIFS({ar("Pipeline General Dwellings")},{ar("Site ID")},{sid})'
    SW[c('Pipeline Student Units')] = f'=SUMIFS({ar("Pipeline Student Units")},{ar("Site ID")},{sid})'
    SW[c('Delivered Units')] = f'=SUMIFS({ar("Delivered Units")},{ar("Site ID")},{sid})'
    SW[c('Not-Counted Applications')] = f'={c("Applications")}-{c("Counted Applications")}'
    SW[c('Low-Confidence Links (unreviewed)')] = f'=COUNTIFS({ar("Site ID")},{sid},{ar("Confidence")},"Low",{ar("Analyst Relationship")},"")'
    SW[c('Stale / Expired (counted)')] = (f'=COUNTIFS({ar("Site ID")},{sid},{ar("Tier Reason")},"Past*")'
                                          f'+COUNTIFS({ar("Site ID")},{sid},{ar("Tier Reason")},"No CIS*")')
    SW[c('Review Required')] = (f'=IF(OR({c("Possible Overlap")}<>"",{c("Low-Confidence Grouping")}="Yes",{c("Low-Confidence Links (unreviewed)")}>0,'
                                f'{c("Stale / Expired (counted)")}>0),IF({c("Site Reviewed")}="Yes","Resolved","Yes"),"No")')
    inp(SW[c('Site Reviewed')])
    inp(SW[c('Site Note')])
    for k in ('Counted Units', 'Pipeline General Dwellings', 'Pipeline Student Units', 'Delivered Units'):
        SW[c(k)].number_format = NUM
dvy = DataValidation(type='list', formula1='"Yes"', allow_blank=True)
SW.add_data_validation(dvy)
dvy.add(f"{CS['Site Reviewed']}2:{CS['Site Reviewed']}{SL}")
sw = {'Origin': 16, 'Site / Project Heading': 42, 'Address': 32, 'Settlement': 18, 'Possible Overlap': 60, 'Site Note': 30,
      'Most Advanced CIS Stage': 18}
for k, col in CS.items():
    SW.column_dimensions[col].width = sw.get(k, 12)
SW.freeze_panes = 'D2'
SW.auto_filter.ref = f'A1:{L(len(scols))}{SL}'
SW.row_dimensions[1].height = 45
rc = CS['Review Required']
SW.conditional_formatting.add(f'{rc}2:{rc}{SL}', FormulaRule(formula=[f'{rc}2="Yes"'], fill=RED))

# ---------------- Review Queue ----------------
cnt_mask = A.rel.isin(COUNTED)
rq = []
for s in S.itertuples():
    ar_ = A[A.site == s.sid]
    iss = []
    low = ar_[(ar_.rel_conf == 'Low') & cnt_mask[ar_.index]]
    if len(low):
        iss.append(f'{len(low)} low-confidence counted application(s): ' + ', '.join(f"{r.ref or r['Project Id']} ({r.rel}, {int(r.Units) if pd.notna(r.Units) else '?'})" for _, r in low.iterrows()))
    med = ar_[(ar_.rel_conf == 'Medium')]
    if len(med):
        iss.append('Auto link to check: ' + '; '.join(f"{r.ref} {r.rel.lower()} {r.rel_ref}" for _, r in med.iterrows()))
    for _, r in ar_[cnt_mask[ar_.index]].iterrows():
        exp = r.exp_final if pd.notna(r.exp_final) else (r['Decision Date'] + pd.DateOffset(years=5) if pd.notna(r['Decision Date']) else None)
        if r['Pipeline Stage'] in ('Permission Granted', 'Commencement Expected') and exp is not None and exp < REPORT_DATE:
            iss.append(f"{r.ref} past expiry ({exp:%d/%m/%Y})")
        elif r['Pipeline Stage'] in ('Permission Granted', 'Commencement Expected') and pd.notna(r['Last Updated']) and (REPORT_DATE - r['Last Updated']).days / 30.44 > 24:
            iss.append(f"{r.ref} no CIS update since {r['Last Updated']:%m/%Y}")
    if s.overlap:
        iss.append('Possible overlap with nearby record')
    if s.lowconf == 'Yes':
        iss.append('Low-confidence grouping in source workbook')
    later = [l for l in ar_.later if l]
    if later:
        iss.append('Later application nearby: ' + later[0])
    if iss:
        units = ar_.loc[cnt_mask[ar_.index], 'Units'].fillna(0).sum()
        rq.append((s.sid, s.heading, s.settlement, s.stage, units, ' | '.join(iss), s.overlap))
rq.sort(key=lambda t: -t[4])
RQ = wb.create_sheet('Review Queue')
rcols = ['#', 'Site ID', 'Site / Project Heading', 'Settlement', 'Most Advanced Stage', 'Counted Units (live)',
         'Issues (at build)', 'Overlap Detail', 'Review Status (live)']
header(RQ, 1, rcols)
for i, t in enumerate(rq, start=2):
    RQ.cell(row=i, column=1, value=i - 1)
    for j, v in enumerate([t[0], t[1], t[2], t[3]], start=2):
        RQ.cell(row=i, column=j, value=nz(v))
    look = lambda k: f'=IFERROR(INDEX(Sites!${CS[k]}$2:${CS[k]}${SL},MATCH($B{i},Sites!$A$2:$A${SL},0)),"")'
    RQ.cell(row=i, column=6, value=look('Counted Units')).number_format = NUM
    RQ.cell(row=i, column=7, value=t[5])
    RQ.cell(row=i, column=8, value=nz(t[6]))
    RQ.cell(row=i, column=9, value=look('Review Required'))
    for col in (7, 8):
        RQ.cell(row=i, column=col).alignment = Alignment(wrap_text=True, vertical='top')
for col, w in zip('ABCDEFGHI', [5, 11, 40, 18, 18, 11, 80, 50, 12]):
    RQ.column_dimensions[col].width = w
RQ.freeze_panes = 'C2'
RQ.auto_filter.ref = f'A1:I{len(rq) + 1}'
RQ.row_dimensions[1].height = 30

# ---------------- Register Not in CIS ----------------
abpnums = {r[6:] for r in CISREFS if r.startswith('ABPREF')}
cand = reg[reg.ApplicationType.isin(['PERMISSION', 'OUTLINE PERMISSION', 'PERMISSION CONSEQUENT'])
           & (np.maximum(reg.NumResidentialUnits.fillna(0), reg.DevelopmentDescription.map(lambda t: desc_units(t) or 0)) >= 10) & (~reg.ref.isin(CISREFS)) & (~reg.ref.str[2:].isin(abpnums))].copy()
cand['out'] = cand.apply(outcome, axis=1)
cand = cand[cand.out.isin(['Pending', 'Under appeal']) |
            (cand.out.isin(['Granted', 'Granted on appeal']) & (cand.ExpiryDate >= REPORT_DATE))]
cand['units_best'] = np.maximum(cand.NumResidentialUnits.fillna(0), cand.DevelopmentDescription.map(lambda t: desc_units(t) or 0)).astype(int)
# Part 8 files held by CIS under a CIS project number (matched by location in the study) are not missing
_p8lab = A[A.ref.str.contains(r'part\s*8|PT8', case=False, na=False)][['Latitude', 'Longitude']].dropna()
_isp8 = cand.ref.str.fullmatch(r'\d{2}8\d{3}')
_nearlab = [bool(len(_p8lab)) and bool((dist(_p8lab.Latitude, _p8lab.Longitude, la, lo) < 60).any()) if pd.notna(la) else False for la, lo in zip(cand.lat, cand.lon)]
cand = cand[~_isp8]  # Part 8s are reconciled in the delivery study
_mod = cand.DevelopmentDescription.fillna('').str.strip().str.lower().str.match(r'(?:the )?(?:minor )?(?:modification|modifying|alteration|amendment|revision|change of house type)')
cand = cand[~_mod]  # amendments to permissions already tracked are not new schemes
siteref = A[A.site != ''][['site', 'Latitude', 'Longitude']].dropna()
NW = wb.create_sheet('Register Not in CIS')
ncols = ['Register Ref', 'Received', 'Outcome', 'Residential Units', 'Expiry', 'Address', 'Description',
         'Nearest CIS Site', 'Distance (m)', 'Register Link', 'Analyst Note']
header(NW, 1, ncols, GFILL)
for i, (_, r) in enumerate(cand.sort_values('units_best', ascending=False).iterrows(), start=2):
    dd = dist(siteref.Latitude, siteref.Longitude, r.lat, r.lon)
    j = dd.idxmin()
    row = [r.ref, r.ReceivedDate, r.out, int(r.units_best), r.ExpiryDate, r.DevelopmentAddress,
           str(r.DevelopmentDescription)[:300], siteref.loc[j, 'site'], int(dd[j]), r.LinkAppDetails]
    for k, v in enumerate(row, 1):
        NW.cell(row=i, column=k, value=nz(v))
    NW.cell(row=i, column=2).number_format = DATE
    NW.cell(row=i, column=5).number_format = DATE
    if r.LinkAppDetails:
        NW.cell(row=i, column=10).hyperlink = r.LinkAppDetails
        NW.cell(row=i, column=10).font = Font(name=F, size=10, color='0563C1', underline='single')
    inp(NW.cell(row=i, column=11))
for col, w in zip('ABCDEFGHIJK', [12, 11, 16, 10, 11, 40, 60, 12, 10, 30, 30]):
    NW.column_dimensions[col].width = w
NW.freeze_panes = 'B2'
NW.auto_filter.ref = f'A1:K{len(cand) + 1}'
NW.row_dimensions[1].height = 30

# ---------------- Summary ----------------
SU = wb.create_sheet('Summary', 1)
SU['A1'] = 'Limerick Residential Pipeline – Summary'
SU['A1'].font = Font(name=F, bold=True, size=14)
SU['A2'] = '="Report date "&TEXT(Method!$B$4,"dd/mm/yyyy")&". Counted applications only (see Relationship). Delivered and Excluded kept out of pipeline totals."'
R = 4
tiers = ['Committed', 'Committed Pre-Construction', 'Advanced Pre-Construction', 'Likely', 'Potential',
         'Application Pipeline', 'Constrained / uncertain']


def sec(t):
    global R
    SU.cell(row=R, column=1, value=t).font = Font(name=F, bold=True, size=12)
    R += 1


def put(row, col, v, fmt=NUM, bold=False):
    cell = SU.cell(row=row, column=col, value=v)
    cell.number_format = fmt
    if bold:
        cell.font = Font(name=F, bold=True)
    return cell


sec('Headline pipeline')
H = {}
for k, f in [('Pipeline general dwellings (units remaining)', f'=SUM({ar("Pipeline General Dwellings")})'),
             ('Pipeline student units', f'=SUM({ar("Pipeline Student Units")})'),
             ('Pipeline student bedspaces', f'=SUM({ar("Pipeline Student Bedspaces")})'),
             ('Of which firm (Committed, Committed Pre-Construction, Advanced Pre-Construction)',
              '=' + '+'.join(f'SUMIFS({ar("Pipeline General Dwellings")},{ar("Reporting Tier")},"{t}")' for t in tiers[:3])),
             ('Delivered units (Complete applications + stated completions)', f'=SUM({ar("Delivered Units")})'),
             ('Counted applications', f'=COUNTIF({ar("Counts?")},"Yes")'),
             ('Applications not counted (superseded, amendments, extensions, not live, excluded)', f'=COUNTIF({ar("Counts?")},"No")')]:
    SU.cell(row=R, column=1, value=k)
    put(R, 2, f)
    H[k] = R
    R += 1
R += 1
sec('Reconciliation to the previous workbook')
SU.cell(row=R, column=1, value='Previous headline units (Pipeline Summary tab, application-level, all stages)')
c = put(R, 2, 11402)
c.font = Font(name=F, color='0000FF')
SU.cell(row=R, column=3, value='Source: Pipeline Summary tab of Limerick_Residential_Pipeline_Full_Build.xlsx (sum of Units).')
prev = R
R += 1
SU.cell(row=R, column=1, value='Units on applications not counted (primacy / not live)')
put(R, 2, f'=SUMIFS({ar("Units")},{ar("Counts?")},"No",{ar("Included in Pipeline")},"Yes")')
R += 1
SU.cell(row=R, column=1, value='Units on counted applications reported as Excluded by analyst')
put(R, 2, f'=SUMIFS({ar("Counted Units")},{ar("Reporting Tier")},"Excluded")')
R += 1
SU.cell(row=R, column=1, value='New total: pipeline general + student units + delivered')
put(R, 2, f'=B{H["Pipeline general dwellings (units remaining)"]}+B{H["Pipeline student units"]}+B{H["Delivered units (Complete applications + stated completions)"]}', bold=True)
R += 1
SU.cell(row=R, column=1, value='Difference to previous headline')
put(R, 2, f'=B{R - 1}-B{prev}')
R += 2

sec('Pipeline by reporting tier')
th = ['Reporting Tier', 'Applications', 'General Dwellings', 'Student Units', 'Student Bedspaces', 'Known Bedroom Mix', 'Mix Coverage']
header(SU, R, th)
R += 1
t0 = R
for t in tiers:
    SU.cell(row=R, column=1, value=t)
    put(R, 2, f'=COUNTIFS({ar("Reporting Tier")},A{R})')
    put(R, 3, f'=SUMIFS({ar("Pipeline General Dwellings")},{ar("Reporting Tier")},A{R})')
    put(R, 4, f'=SUMIFS({ar("Pipeline Student Units")},{ar("Reporting Tier")},A{R})')
    put(R, 5, f'=SUMIFS({ar("Pipeline Student Bedspaces")},{ar("Reporting Tier")},A{R})')
    put(R, 6, f'=SUMIFS({ar("Mix Total")},{ar("Reporting Tier")},A{R},{ar("Dwelling Type")},"General")')
    put(R, 7, f'=IF(SUMIFS({ar("Counted Units")},{ar("Reporting Tier")},A{R},{ar("Dwelling Type")},"General")=0,"",'
              f'MIN(1,F{R}/SUMIFS({ar("Counted Units")},{ar("Reporting Tier")},A{R},{ar("Dwelling Type")},"General")))', '0%')
    R += 1
SU.cell(row=R, column=1, value='Pipeline total').font = Font(name=F, bold=True)
for col in 'BCDEF':
    put(R, ' BCDEF'.index(col) + 1, f'=SUM({col}{t0}:{col}{R - 1})', bold=True)
put(R, 7, '=IF(' + '+'.join(f'SUMIFS({ar("Counted Units")},{ar("Reporting Tier")},A{k},{ar("Dwelling Type")},"General")' for k in range(t0, R)) +
    f'=0,"",MIN(1,F{R}/(' + '+'.join(f'SUMIFS({ar("Counted Units")},{ar("Reporting Tier")},A{k},{ar("Dwelling Type")},"General")' for k in range(t0, R)) + ')))', '0%', bold=True)
pipe_total_row = R
R += 1
for t in ['Delivered', 'Excluded']:
    SU.cell(row=R, column=1, value=t)
    put(R, 2, f'=COUNTIFS({ar("Reporting Tier")},A{R})')
    col = 'Delivered Units' if t == 'Delivered' else 'Counted Units'
    put(R, 3, f'=SUMIFS({ar(col)},{ar("Reporting Tier")},A{R},{ar("Dwelling Type")},"General")')
    R += 1
SU.cell(row=R, column=1, value='Mix coverage = counted general units with a known bedroom split, as a share of counted general units in the tier.').font = Font(name=F, italic=True, size=9)
R += 2

sec('Pipeline bedroom mix (general dwellings, counted applications with known mix)')
bh = ['Bedrooms', 'Units', 'Share', 'Of which from CIS / KPMG', 'Of which from register description']
header(SU, R, bh)
R += 1
b0 = R
for b in BEDCOLS:
    SU.cell(row=R, column=1, value=b)
    base = f'{ar(b)},{ar("Reporting Tier")},"<>Not counted",{ar("Reporting Tier")},"<>Delivered",{ar("Reporting Tier")},"<>Excluded",{ar("Dwelling Type")},"General"'
    put(R, 2, f'=SUMIFS({base})')
    put(R, 4, f'=SUMIFS({base},{ar("Mix Source")},"CIS / KPMG")')
    put(R, 5, f'=SUMIFS({base},{ar("Mix Source")},"Register description")')
    R += 1
SU.cell(row=R, column=1, value='Total with known mix').font = Font(name=F, bold=True)
for col in 'BDE':
    put(R, ' ABCDE'.index(col), f'=SUM({col}{b0}:{col}{R - 1})', bold=True)
for k in range(b0, R):
    put(k, 3, f'=IF($B${R}=0,"",B{k}/$B${R})', '0%')
R += 1
SU.cell(row=R, column=1, value='Counted general units in pipeline (all, incl. unknown mix)')
put(R, 2, f'=SUMIFS({ar("Counted Units")},{ar("Reporting Tier")},"<>Not counted",{ar("Reporting Tier")},"<>Delivered",{ar("Reporting Tier")},"<>Excluded",{ar("Dwelling Type")},"General")')
R += 1
SU.cell(row=R, column=1, value='Mix coverage')
put(R, 2, f'=IF(B{R - 1}=0,"",B{R - 2}/B{R - 1})', '0%')
R += 1
SU.cell(row=R, column=1, value='Shares describe permitted units on applications with a known mix, not the whole pipeline. Treat as indicative where coverage is low.').font = Font(name=F, italic=True, size=9)
R += 2

sec('Application relationships (primacy)')
rh = ['Final Relationship', 'Applications', 'Units (CIS)', 'Counted?', 'Unreviewed Low / Medium confidence']
header(SU, R, rh)
R += 1
for rl in RELS:
    SU.cell(row=R, column=1, value=rl)
    put(R, 2, f'=COUNTIF({ar("Final Relationship")},A{R})')
    put(R, 3, f'=SUMIFS({ar("Units")},{ar("Final Relationship")},A{R})')
    SU.cell(row=R, column=4, value='Yes' if rl in COUNTED else 'No')
    put(R, 5, f'=COUNTIFS({ar("Final Relationship")},A{R},{ar("Confidence")},"Low",{ar("Analyst Relationship")},"")'
              f'+COUNTIFS({ar("Final Relationship")},A{R},{ar("Confidence")},"Medium",{ar("Analyst Relationship")},"")')
    R += 1
R += 1

sec('Pipeline by settlement (general dwellings)')
setts = sorted(S.settlement.unique(), key=lambda s: (s == 'Other County Limerick', s))
sh = ['Settlement', 'Applications', 'Firm', 'Likely + Potential', 'Application Pipeline', 'Constrained / uncertain', 'Total']
header(SU, R, sh)
R += 1
s0 = R
for st in setts:
    SU.cell(row=R, column=1, value=st)
    put(R, 2, f'=COUNTIFS({ar("Settlement")},A{R},{ar("Pipeline General Dwellings")},">0")')
    g = lambda tl: '=' + '+'.join(f'SUMIFS({ar("Pipeline General Dwellings")},{ar("Settlement")},$A{R},{ar("Reporting Tier")},"{t}")' for t in tl)
    put(R, 3, g(tiers[:3]))
    put(R, 4, g(['Likely', 'Potential']))
    put(R, 5, g(['Application Pipeline']))
    put(R, 6, g(['Constrained / uncertain']))
    put(R, 7, f'=SUM(C{R}:F{R})')
    R += 1
SU.cell(row=R, column=1, value='Total').font = Font(name=F, bold=True)
for k, col in enumerate('BCDEFG', start=2):
    put(R, k, f'=SUM({col}{s0}:{col}{R - 1})', bold=True)
R += 1
SU.cell(row=R, column=1, value='Settlement for upstream records is assigned from address text; check before using at settlement level.').font = Font(name=F, italic=True, size=9)
R += 2

sec('Register validation (applications)')
header(SU, R, ['Check Result', 'Applications'])
R += 1
v0 = R
for res in ['Consistent', 'Consistent – decision date differs', 'Stage out of date – now granted',
            'Check – register still pending', 'Check – location differs', 'Conflict – not live in register',
            'Conflict – register not live but CIS says built', 'Unverified – not in register (Part 8 / other)',
            'Unverified – no reference', 'Excluded (non-residential)']:
    SU.cell(row=R, column=1, value=res)
    put(R, 2, f'=COUNTIF({ar("Check Result")},A{R})')
    R += 1
SU.cell(row=R, column=1, value='Total').font = Font(name=F, bold=True)
put(R, 2, f'=SUM(B{v0}:B{R - 1})', bold=True)
R += 2
sec('Review status')
srng = lambda k: f"Sites!${CS[k]}$2:${CS[k]}${SL}"
for k, f in [('Sites needing review', f'=COUNTIF({srng("Review Required")},"Yes")'),
             ('Sites marked reviewed', f'=COUNTIF({srng("Review Required")},"Resolved")'),
             ('Counted units on sites needing review', f'=SUMIFS({srng("Counted Units")},{srng("Review Required")},"Yes")'),
             ('Unreviewed low-confidence counted applications', f'=COUNTIFS({ar("Confidence")},"Low",{ar("Counts?")},"Yes",{ar("Analyst Relationship")},"")')]:
    SU.cell(row=R, column=1, value=k)
    put(R, 2, f)
    R += 1
SU.column_dimensions['A'].width = 66
for col in 'BCDEFG':
    SU.column_dimensions[col].width = 17

# fonts
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            f = c.font
            if f.name != F:
                c.font = Font(name=F, size=10 if (not f.size or f.size == 11) else f.size, bold=f.b, italic=f.i,
                              color=f.color, underline=f.u)
wb._sheets = [wb['Method'], wb['Summary'], wb['Sites'], wb['Review Queue'], wb['Applications'], wb['Register Not in CIS']]
wb.save(OUT)
print('apps', N, 'sites', len(S), 'queue', len(rq), 'reg-not-in-cis', len(cand))
print(A.rel.value_counts().to_string())
print(A.src.value_counts().to_string())
